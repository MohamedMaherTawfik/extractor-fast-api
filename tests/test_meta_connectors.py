from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
import re

import httpx
import pytest
from openpyxl import load_workbook

from backend.core.config import Settings
from backend.creator_discovery.connectors import FacebookMetaConnector, InstagramMetaConnector
from backend.creator_discovery.connectors import CreatorDiscoveryConnectorRegistry
from backend.creator_discovery.meta import MetaApiClient, MetaApiError
from backend.creator_discovery.normalization import normalize_discovery_input
from backend.creator_discovery.quality import data_completeness
from backend.creator_discovery.service import CreatorDiscoveryService
from backend.db.session import session_scope
from backend.schemas.creator_discovery import (
    CandidateDecisionRequest,
    CreatorExportRequest,
    CreatorDiscoveryRunRequest,
    CreatorRefreshRequest,
    DiscoveryPlatform,
)


def _settings(**overrides):
    values = {
        "meta_app_id": "app-123",
        "meta_app_secret": "server-secret",
        "meta_access_token": "operator-token-A91B",
        "meta_graph_api_version": "v26.0",
        "meta_facebook_page_id": "page-1",
        "meta_instagram_account_id": "ig-1",
    }
    values.update(overrides)
    return Settings(**values)


def _json(data, status=200, headers=None):
    return httpx.Response(status, json=data, headers=headers)


def _valid_handler(request: httpx.Request) -> httpx.Response:
    path = request.url.path
    if path.endswith("/debug_token"):
        return _json({"data": {
            "is_valid": True,
            "app_id": "app-123",
            "expires_at": int(datetime(2030, 1, 1, tzinfo=UTC).timestamp()),
            "scopes": ["instagram_basic", "pages_read_engagement", "pages_show_list"],
        }})
    if path.endswith("/ig-1/media"):
        return _json({"data": [{
            "id": "media-1", "caption": "Beauty tips #makeup", "media_type": "IMAGE",
            "media_url": "https://cdn.example/media.jpg", "permalink": "https://www.instagram.com/p/one/",
            "timestamp": "2026-08-01T12:00:00+0000", "like_count": 120, "comments_count": 7,
        }]})
    if path.endswith("/ig-1"):
        return _json({
            "id": "ig-1", "username": "brand", "name": "Brand Studio", "biography": "Beauty and cosmetics",
            "followers_count": 5000, "follows_count": 50, "media_count": 90,
            "profile_picture_url": "https://cdn.example/avatar.jpg", "website": "https://brand.example",
        })
    if path.endswith("/page-1/posts"):
        return _json({"data": [{
            "id": "page-1_10", "message": "New product #launch", "created_time": "2026-08-02T12:00:00+0000",
            "permalink_url": "https://www.facebook.com/brand/posts/10", "full_picture": "https://cdn.example/post.jpg",
            "likes": {"summary": {"total_count": 45}}, "comments": {"summary": {"total_count": 3}},
            "shares": {"count": 2}, "attachments": {"data": [{"type": "photo", "url": "https://www.facebook.com/10"}]},
        }]})
    if path.endswith("/page-1") or path.endswith("/brand"):
        return _json({
            "id": "page-1", "name": "Brand Studio", "username": "brand",
            "link": "https://www.facebook.com/brand", "about": "Official beauty Page",
            "fan_count": 7000, "followers_count": 7100, "verification_status": "blue_verified",
            "picture": {"data": {"url": "https://cdn.example/page.jpg"}},
        })
    raise AssertionError(f"Unexpected Meta request: {request.url}")


def _client(handler=_valid_handler, **settings):
    return MetaApiClient(
        _settings(**settings),
        transport=httpx.MockTransport(handler),
        sleeper=lambda _: None,
    )


def test_missing_meta_credentials_are_not_configured():
    client = MetaApiClient(Settings(
        meta_app_id=None, meta_app_secret=None, meta_access_token=None,
        meta_facebook_page_id=None, meta_instagram_account_id=None,
    ))
    result = client.validate().public_payload()
    assert result["token_status"] == "NOT_CONFIGURED"
    assert result["instagram"]["status"] == "NOT_CONFIGURED"
    assert result["facebook"]["status"] == "NOT_CONFIGURED"
    assert result["token_hint"] is None


def test_meta_status_api_returns_flags_without_secret_values(api_request):
    response = api_request("GET", "/creator-discovery/meta/status")
    assert response.status_code == 200
    payload = response.json()
    assert payload["token_status"] in {"NOT_CONFIGURED", "READY", "TOKEN_INVALID", "TOKEN_EXPIRED", "API_UNAVAILABLE", "RATE_LIMITED"}
    assert "app_secret" not in payload and "access_token" not in payload
    assert "status" in payload["instagram"] and "status" in payload["facebook"]


@pytest.mark.parametrize(
    ("debug_data", "expected"),
    [
        ({"is_valid": False, "app_id": "app-123"}, "TOKEN_INVALID"),
        ({"is_valid": True, "app_id": "app-123", "expires_at": 1}, "TOKEN_EXPIRED"),
    ],
)
def test_invalid_and_expired_tokens(debug_data, expected):
    def handler(request):
        return _json({"data": debug_data})

    result = _client(handler).validate()
    assert result.token_status == expected
    assert result.instagram.status == expected
    assert "operator-token-A91B" not in result.token_message


def test_missing_permissions_are_reported_per_platform():
    def handler(request):
        return _json({"data": {"is_valid": True, "app_id": "app-123", "scopes": ["pages_show_list"]}})

    result = _client(handler).validate()
    assert result.token_status == "READY"
    assert result.instagram.status == "PERMISSION_MISSING"
    assert result.facebook.status == "PERMISSION_MISSING"
    assert "instagram_basic" in result.instagram.missing_permissions


def test_instagram_profile_content_normalization_and_provenance():
    connector = InstagramMetaConnector(_client())
    outcome = connector.discover(normalize_discovery_input("https://www.instagram.com/brand/"))
    assert outcome.status == "FOUND"
    candidate = outcome.candidates[0]
    assert candidate.username == "brand" and candidate.followers == 5000
    assert candidate.discovery_source == "meta_api"
    assert candidate.provenance[0]["api_object_id"] == "ig-1"
    assert candidate.data_completeness > 50
    samples = connector.collect_recent(candidate, 10)
    assert len(samples) == 1 and samples[0].likes == 120
    assert samples[0].platform_content_id == "media-1"
    assert samples[0].provenance[0]["source"] == "meta_api"


def test_facebook_profile_content_normalization_and_provenance():
    connector = FacebookMetaConnector(_client())
    outcome = connector.discover(normalize_discovery_input("https://www.facebook.com/brand/"))
    assert outcome.status == "FOUND"
    candidate = outcome.candidates[0]
    assert candidate.source_id == "page-1" and candidate.followers == 7100
    assert candidate.verified is True
    samples = connector.collect_recent(candidate, 10)
    assert len(samples) == 1
    assert (samples[0].likes, samples[0].comments, samples[0].shares) == (45, 3, 2)
    assert samples[0].content_url.endswith("/posts/10")


def test_rate_limit_uses_bounded_retries_and_safe_status():
    attempts = []

    def handler(request):
        attempts.append(request)
        return _json({"error": {"code": 4, "message": "Application request limit reached"}}, 429, {"Retry-After": "0"})

    client = _client(handler)
    with pytest.raises(MetaApiError) as raised:
        client.get_facebook_page("page-1")
    assert raised.value.status == "RATE_LIMITED"
    assert len(attempts) == 3
    assert "operator-token-A91B" not in str(raised.value)


def test_data_completeness_is_deterministic_and_not_match_confidence():
    connector = InstagramMetaConnector(_client())
    candidate = connector.discover(normalize_discovery_input("https://www.instagram.com/brand")).candidates[0]
    first = data_completeness(candidate)
    second = data_completeness(candidate)
    assert first == second
    assert first != 100.0  # analysis/content evidence are not inferred from identity matching


def test_meta_refresh_updates_existing_profile_without_duplication():
    state = {"followers": 5000}

    def handler(request):
        response = _valid_handler(request)
        if request.url.path.endswith("/ig-1"):
            body = response.json()
            body["followers_count"] = state["followers"]
            return _json(body)
        return response

    meta = _client(handler)
    registry = CreatorDiscoveryConnectorRegistry(connectors=[InstagramMetaConnector(meta)])
    with session_scope() as session:
        service = CreatorDiscoveryService(session, connectors=registry)
        run = service.create_run(CreatorDiscoveryRunRequest(
            inputs=["https://www.instagram.com/brand"],
            platforms=[DiscoveryPlatform.INSTAGRAM],
            content_sample_size=10,
        ))
        candidate = service.list_candidates(
            offset=0, limit=10, run_uid=run.run_uid, platform="instagram",
            review_status="PENDING", classification=None,
        )["items"][0]
        created = service.confirm_candidate(
            candidate.candidate_uid,
            CandidateDecisionRequest(action="THIS_IS_THE_ACCOUNT"),
        )
        profile_uid = created["unified_profile"]["profile_uid"]
        state["followers"] = 6200
        refreshed = service.refresh_profile(profile_uid, CreatorRefreshRequest(
            platforms=[DiscoveryPlatform.INSTAGRAM], content_sample_size=10,
        ))
        assert refreshed["unified_profile"]["profile_uid"] == profile_uid
        assert len(refreshed["platform_accounts"]) == 1
        assert refreshed["platform_accounts"][0]["followers"] == 6200
        assert len(refreshed["content_samples"]) == 1
        assert refreshed["refresh_status"]["instagram"] == "FOUND"


def test_five_instagram_urls_collect_analyze_export_and_repair_identity_only_account():
    usernames = [f"beauty_creator_{index}" for index in range(1, 6)]

    def invalid_handler(request):
        if request.url.path.endswith("/debug_token"):
            return _json({"data": {"is_valid": False, "app_id": "app-123"}})
        raise AssertionError(f"Unexpected Meta request: {request.url}")

    def valid_bulk_handler(request):
        path = request.url.path
        fields = request.url.params.get("fields", "")
        if path.endswith("/debug_token"):
            return _json({"data": {
                "is_valid": True,
                "app_id": "app-123",
                "expires_at": int(datetime(2030, 1, 1, tzinfo=UTC).timestamp()),
                "scopes": ["instagram_basic", "pages_read_engagement", "pages_show_list"],
            }})
        if path.endswith("/ig-1") and fields == "id,username,name":
            return _json({"id": "ig-1", "username": "connected_account", "name": "Connected Account"})
        if path.endswith("/ig-1") and fields.startswith("id,username,name,biography"):
            return _json({"id": "ig-1", "username": "connected_account", "name": "Connected Account"})
        if path.endswith("/ig-1") and "business_discovery.username" in fields:
            match = re.search(r"business_discovery\.username\(([^)]+)\)", fields)
            assert match is not None
            username = match.group(1)
            if "media.limit" in fields:
                return _json({"business_discovery": {"media": {"data": [
                    {
                        "id": f"{username}-media-{index}",
                        "caption": f"Skincare tutorial tips — follow for more #skincare #beauty {index}",
                        "media_type": "IMAGE",
                        "media_url": f"https://cdn.example/{username}-{index}.jpg",
                        "permalink": f"https://www.instagram.com/p/{username}-{index}/",
                        "timestamp": f"2026-08-0{index}T12:00:00+0000",
                        "like_count": 6500,
                        "comments_count": 150,
                    }
                    for index in range(1, 4)
                ]}}})
            creator_index = int(username.rsplit("_", 1)[1])
            return _json({"business_discovery": {
                "id": f"ig-{creator_index + 10}",
                "username": username,
                "name": f"Beauty Creator {creator_index}",
                "biography": "Beauty, skincare tutorials, makeup tips, and product reviews.",
                "followers_count": 100_000 + creator_index,
                "follows_count": 200 + creator_index,
                "media_count": 300 + creator_index,
                "profile_picture_url": f"https://cdn.example/{username}.jpg",
            }})
        raise AssertionError(f"Unexpected Meta request: {request.url}")

    invalid_meta = _client(invalid_handler)
    invalid_registry = CreatorDiscoveryConnectorRegistry(connectors=[InstagramMetaConnector(invalid_meta)])
    valid_meta = _client(valid_bulk_handler, meta_facebook_page_id=None)
    valid_registry = CreatorDiscoveryConnectorRegistry(connectors=[InstagramMetaConnector(valid_meta)])

    with session_scope() as session:
        invalid_service = CreatorDiscoveryService(session, connectors=invalid_registry)
        first_run = invalid_service.create_run(CreatorDiscoveryRunRequest(
            inputs=[f"https://www.instagram.com/{usernames[0]}"],
            platforms=[DiscoveryPlatform.INSTAGRAM],
            mode="bulk",
            auto_process=True,
            analyze_content=True,
        ))
        assert first_run.status == "PARTIAL"
        identity_only = invalid_service.list_profiles(offset=0, limit=10, include_export_fields=True)["items"][0]
        assert identity_only["followers"] is None
        assert identity_only["analysis_status"] == "TOKEN_INVALID"

        service = CreatorDiscoveryService(session, connectors=valid_registry)
        run = service.create_run(CreatorDiscoveryRunRequest(
            inputs=[f"https://www.instagram.com/{username}" for username in usernames],
            platforms=[DiscoveryPlatform.INSTAGRAM],
            mode="bulk",
            auto_process=True,
            analyze_content=True,
            update_existing_profiles=False,
            content_sample_size=10,
        ))
        assert run.status == "COMPLETED", [
            (job.status, job.error, job.checkpoint) for job in run.jobs
        ]
        assert run.processed == 5 and run.matched == 5 and run.review_required == 0

        page = service.list_profiles(offset=0, limit=10, include_export_fields=True)
        assert page["total"] == 5
        for row in page["items"]:
            assert row["followers"] is not None
            assert row["following"] is not None
            assert row["content_count"] is not None
            assert row["bio"]
            assert row["industry"] != "UNKNOWN"
            assert row["niche"] != "UNKNOWN"
            assert row["content_type"] == "IMAGE"
            assert row["content_angle"] != "UNKNOWN"
            assert row["content_style"] != "UNKNOWN"
            assert row["visual_style"] == "Static image-led"
            assert row["cta"] == "subscribe"
            assert row["engagement"] != "UNKNOWN"
            assert row["audience_analysis"] != "UNKNOWN"
            assert row["trend_analysis"] != "UNKNOWN"
            assert row["ai_recommendation"] != "INSUFFICIENT_EVIDENCE"
            assert row["campaign_idea"] != "INSUFFICIENT_EVIDENCE"
            assert row["kpi_impact"] != "UNKNOWN"

            profile = service.repository.get_profile_by_creator_uid(row["creator_uid"])
            assert profile is not None
            assert len(service.repository.samples_for_creator(profile.creator_id)) == 3
            assert len(service.repository.analyses_for_profile(profile.id)) >= 1

        payload, _, _ = service.export(CreatorExportRequest(format="xlsx", run_uid=run.run_uid))
        sheet = load_workbook(BytesIO(payload))["Creator Discovery"]
        assert sheet.max_row == 6
        assert all(sheet.cell(row=index, column=5).value for index in range(2, 7))
        assert all(sheet.cell(row=index, column=10).value != "UNKNOWN" for index in range(2, 7))
        assert all(sheet.cell(row=index, column=21).value != "INSUFFICIENT_EVIDENCE" for index in range(2, 7))
