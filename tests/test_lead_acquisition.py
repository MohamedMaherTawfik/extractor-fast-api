from __future__ import annotations

from io import BytesIO
import sys
from types import ModuleType, SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from backend.db.session import session_scope
from backend.leads.collectors import CollectionContext, OSMGeofabrikCollector, OvertureCollector, SourceRateLimitError
from backend.leads.geography import GeographyService
from backend.leads.normalization import LeadNormalizer, RawLeadRecord
from backend.leads.scoring import LeadScorer
from backend.leads.service import LeadAcquisitionService
from backend.leads.sources import LeadSourceRegistry
from backend.leads.workbook import LeadWorkbookService
from backend.schemas.leads import LeadRunRequest


class FakeRecordBatch:
    def __init__(self, rows: list[dict]) -> None:
        self._rows = rows

    def to_pylist(self) -> list[dict]:
        return self._rows


def install_fake_overture_reader(monkeypatch, rows: list[dict]) -> None:
    module = ModuleType("overturemaps")
    module.record_batch_reader = lambda *args, **kwargs: [FakeRecordBatch(rows)]
    monkeypatch.setitem(sys.modules, "overturemaps", module)
    shapely = ModuleType("shapely")
    shapely.from_wkb = lambda value: SimpleNamespace(x=31.2357, y=30.0444, geom_type="Point")
    monkeypatch.setitem(sys.modules, "shapely", shapely)


def test_numbered_workbook_aliases_and_structured_headers(tmp_path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    definitions = {
        "10_Lead_Segments": (["Category_ID", "Phase", "Tier", "Segment_EN", "Buyer_Type"], ["CAT-001", 1, "A+", "Cosmetics Importers", "B2B"]),
        "11_Keyword_Master": (["Keyword_ID", "Category_ID", "Keyword", "Lang", "Status"], ["KW-001", "CAT-001", "cosmetics importer", "EN", "ENABLED"]),
        "12_Egypt_Coverage": (["Gov_ID", "Governorate_EN", "Governorate_AR"], ["GOV-01", "Cairo", "القاهرة"]),
        "13_Source_Registry": (["Source_ID", "Source"], ["SRC-001", "Overture"]),
        "14_Query_Matrix": (["Query_ID", "Enabled", "Keyword_ID"], ["Q-001", True, "KW-001"]),
        "19_Run_Config": (["Config_ID", "Key", "Value"], ["CFG-001", "country_code", "EG"]),
    }
    for name, (headers, values) in definitions.items():
        sheet = workbook.create_sheet(name)
        sheet.append(["Title"])
        sheet.append(["Description"])
        sheet.append(headers)
        sheet.append(values)
    target = tmp_path / "numbered.xlsx"
    workbook.save(target)

    service = LeadWorkbookService()
    preview = service.preview(target)
    assert preview["status"] == "VALID"
    assert preview["active_command_source"] == "WORKBOOK"
    assert preview["resolved_sheets"]["Lead Segments"] == "10_Lead_Segments"
    assert preview["sheet_counts"]["Keyword Master"] == 1
    assert preview["query_matrix_rows"] == preview["enabled_query_jobs"] == 1
    assert service.rows("Keyword Master", target)[0]["Keyword"] == "cosmetics importer"


def test_actual_command_workbook_counts_when_present() -> None:
    service = LeadWorkbookService()
    if service.find() is None:
        pytest.skip("Runtime command workbook is not present")
    preview = service.preview()
    assert preview["status"] == "VALID"
    assert preview["active_command_source"] == "WORKBOOK"
    assert preview["sheet_counts"]["Keyword Master"] == 431
    assert preview["sheet_counts"]["Lead Segments"] == 50
    assert preview["sheet_counts"]["Egypt Coverage"] == 27
    assert preview["query_matrix_rows"] == 11_475
    assert preview["enabled_query_jobs"] == 11_475


def test_source_registry_and_full_egypt_control_catalog() -> None:
    registry = LeadSourceRegistry()
    statuses = {item.config["source_uid"]: item.status for item in registry.describe_all()}
    assert statuses["SRC_OVERTURE"] == "READY"
    assert statuses["SRC_OSM_GEOFABRIK"] == "READY"
    assert statuses["SRC_GOOGLE_PLACES"] == "NOT_CONFIGURED"
    assert statuses["SRC_FOURSQUARE"] == "NOT_CONFIGURED"
    assert statuses["SRC_DIRECTORY"] == "DISABLED"
    with session_scope() as session:
        options = LeadAcquisitionService(session).control_options()
        assert len(options["governorates"]) == 27
        if options["workbook"]["status"] == "VALID":
            assert options["active_command_source"] == "WORKBOOK"
            assert len(options["segments"]) == 50
            assert options["keywords"]["count"] == 431
            assert options["query_matrix_rows"] == options["enabled_query_jobs"] == 11_475
        else:
            assert len(options["segments"]) >= 37
            assert options["keywords"]["count"] >= 70
        assert options["workbook"]["placement"].startswith("data/imports/lead_control/")


def test_geography_adaptive_tiles_and_point_resolution() -> None:
    geography = GeographyService()
    cairo = geography.adaptive_tiles(["cairo"])
    new_valley = geography.adaptive_tiles(["new_valley"])
    assert cairo and new_valley
    assert all(len(tile.polygon) == 5 for tile in cairo)
    assert geography.governorate_for_point(31.2357, 30.0444) == "cairo"
    assert len(geography.list_governorates()) == 27


def test_normalization_is_conservative_for_arabic_phone_domain_and_category() -> None:
    normalizer = LeadNormalizer()
    raw = RawLeadRecord(
        source_uid="SRC_OVERTURE",
        source_record_id="one",
        business_name="صيدلية النُّور",
        category_raw="pharmacy",
        address="١٢ شارع التحرير",
        latitude=30.0444,
        longitude=31.2357,
        phone="٠١٠١٢٣٤٥٦٧٨",
        website="https://www.Example.COM/path",
    )
    result = normalizer.normalize(raw)
    assert result["business_name_normalized"] == "صيدلية النور"
    assert result["phone_normalized"] == "+201012345678"
    assert result["domain_normalized"] == "example.com"
    assert result["category_id"] == "pharmacies"
    assert result["governorate"] == "cairo"
    assert normalizer.normalize_phone("12345") is None


def test_scoring_is_config_driven_and_classified() -> None:
    values = {
        "segment_tier": 1, "buyer_type": "WHOLESALER", "phone_normalized": "+201000000000",
        "email": "sales@example.com", "website": "https://example.com", "social_links": ["https://example.com/social"],
        "governorate": "cairo", "last_verified_at": None, "source_last_seen_at": __import__("datetime").datetime.now(__import__("datetime").UTC),
    }
    score, fit_class, explanation = LeadScorer().score(values, source_count=2)
    assert score >= 85
    assert fit_class == "A+"
    assert set(explanation["components"]) == set(explanation["weights"])


def test_dedupe_merges_cross_source_but_preserves_distinct_branches() -> None:
    with session_scope() as session:
        service = LeadAcquisitionService(session)
        run = service.create_run(LeadRunRequest(name="dedupe", sources=["SRC_OVERTURE"], governorates=["cairo"], segments=["pharmacies"], execute=False))
        job = run.jobs[0]
        first = RawLeadRecord(source_uid="SRC_OVERTURE", source_record_id="ov-1", business_name="Al Noor Pharmacy", category_raw="pharmacy", address="12 Tahrir St", latitude=30.0444, longitude=31.2357, phone="01012345678")
        second = RawLeadRecord(source_uid="SRC_OSM_GEOFABRIK", source_record_id="node/2", business_name="Al Noor Pharmacy", category_raw="pharmacy", address="12 Tahrir St", latitude=30.04441, longitude=31.23571, phone="01012345678")
        branch = RawLeadRecord(source_uid="SRC_OSM_GEOFABRIK", source_record_id="node/3", business_name="Al Noor Pharmacy", category_raw="pharmacy", address="90 Other St", latitude=30.10, longitude=31.30, phone="01012345678")
        assert service._persist_record(run, job, first, "data/lead_acquisition/raw/a.jsonl") == "UNIQUE"
        assert service._persist_record(run, job, first, "data/lead_acquisition/raw/a2.jsonl") == "SAME_SOURCE_DUPLICATE"
        assert service._persist_record(run, job, second, "data/lead_acquisition/raw/b.jsonl") == "CROSS_SOURCE_MERGED"
        assert service._persist_record(run, job, branch, "data/lead_acquisition/raw/c.jsonl") == "UNIQUE"
        page = service.list_leads(offset=0, limit=50, fit_class=None, governorate=None, category=None, source=None, verified_only=False, query=None)
        assert page["total"] == 2
        detail = service.get_lead(page["items"][0]["lead_uid"])
        assert detail["source_records"]


def test_dedupe_domain_signal_merges_and_keeps_source_provenance() -> None:
    with session_scope() as session:
        service = LeadAcquisitionService(session)
        run = service.create_run(LeadRunRequest(name="domain dedupe", sources=["SRC_OVERTURE"], governorates=["cairo"], segments=["cosmetics_suppliers"], execute=False))
        job = run.jobs[0]
        first = RawLeadRecord(source_uid="SRC_OVERTURE", source_record_id="domain-1", business_name="Beauty Supply Egypt", category_raw="beauty_supply_store", website="https://www.beauty.example/", latitude=30.04, longitude=31.23)
        second = RawLeadRecord(source_uid="SRC_OSM_GEOFABRIK", source_record_id="node/domain-2", business_name="Beauty Supply EG", category_raw="beauty_supply_store", website="https://beauty.example/contact", latitude=30.0401, longitude=31.2301)
        assert service._persist_record(run, job, first, "data/lead_acquisition/raw/domain-1.jsonl") == "UNIQUE"
        assert service._persist_record(run, job, second, "data/lead_acquisition/raw/domain-2.jsonl") == "CROSS_SOURCE_MERGED"
        page = service.list_leads(offset=0, limit=10, fit_class=None, governorate=None, category=None, source=None, verified_only=False, query=None)
        assert page["total"] == 1
        detail = service.get_lead(page["items"][0]["lead_uid"])
        assert {item["source_uid"] for item in detail["source_records"]} == {"SRC_OVERTURE", "SRC_OSM_GEOFABRIK"}


def test_run_dry_plan_checkpoint_and_controls() -> None:
    with session_scope() as session:
        service = LeadAcquisitionService(session)
        dry = service.create_run(LeadRunRequest(name="dry", sources=["SRC_OVERTURE"], governorates=["cairo"], segments=["pharmacies"], dry_run=True, execute=False))
        assert dry.dry_run and dry.planned_jobs > 0 and dry.keyword_count >= 2
        run = service.create_run(LeadRunRequest(name="control", sources=["SRC_OVERTURE"], governorates=["cairo"], segments=["pharmacies"], execute=False))
        run.jobs[0].checkpoint = {"raw_position": 75}
        service.pause(run.run_uid)
        assert service.resume(run.run_uid).status == "PENDING"
        loaded = service.get_run(run.run_uid)
        assert loaded["jobs"][0]["checkpoint"]["raw_position"] == 75
        assert service.cancel(run.run_uid).status == "CANCELLED"


def test_overture_collector_streams_filtered_sample(monkeypatch) -> None:
    row = {
        "id": "real-shape-sample", "geometry": b"point",
        "categories": {"primary": "pharmacy", "alternate": []}, "confidence": 0.92,
        "websites": ["https://example.test"], "emails": [], "socials": [], "phones": ["+201000000000"],
        "brand": None, "addresses": [{"freeform": "Cairo", "locality": "Cairo", "postcode": None, "region": "Cairo", "country": "EG"}],
        "names": {"primary": "Sample Pharmacy", "common": None, "rules": None}, "sources": [],
        "operating_status": "open", "basic_category": "business", "taxonomy": None, "version": 1,
        "bbox": {"xmin": 31.2357, "xmax": 31.2357, "ymin": 30.0444, "ymax": 30.0444},
    }
    install_fake_overture_reader(monkeypatch, [row])
    monkeypatch.setattr(OvertureCollector, "discover_latest_release", lambda self: "test-release")
    records = []
    context = CollectionContext(source_uid="SRC_OVERTURE", run_uid="r", job_uid="j", governorate="cairo", tile={"bbox": [31.2, 30.0, 31.3, 30.1]}, segment_ids=["pharmacies"], keywords=["pharmacy"], max_records=10)
    result = OvertureCollector().collect(context, lambda record, position: records.append(record) or True)
    assert result["emitted"] == 1
    assert records[0].source_record_id == "real-shape-sample"


def test_overture_checkpoint_resume_skips_already_processed_rows(monkeypatch) -> None:
    base = {
        "geometry": b"point", "categories": {"primary": "pharmacy", "alternate": []},
        "confidence": 0.9, "websites": [], "emails": [], "socials": [], "phones": [], "brand": None,
        "addresses": [], "names": {"primary": "Pharmacy", "common": None, "rules": None}, "sources": [],
        "operating_status": "open", "basic_category": "business", "taxonomy": None, "version": 1,
        "bbox": {"xmin": 31.2357, "xmax": 31.2357, "ymin": 30.0444, "ymax": 30.0444},
    }
    install_fake_overture_reader(monkeypatch, [{**base, "id": "already-done"}, {**base, "id": "resume-here"}])
    monkeypatch.setattr(OvertureCollector, "discover_latest_release", lambda self: "test-release")
    records = []
    context = CollectionContext(source_uid="SRC_OVERTURE", run_uid="r", job_uid="j", governorate="cairo", tile={"bbox": [31.2, 30.0, 31.3, 30.1]}, segment_ids=["pharmacies"], keywords=["pharmacy"], checkpoint={"raw_position": 1})
    OvertureCollector().collect(context, lambda record, position: records.append(record) or True)
    assert [record.source_record_id for record in records] == ["resume-here"]


def test_scoring_produces_all_four_classes_deterministically() -> None:
    from datetime import UTC, datetime, timedelta

    now = datetime.now(UTC)
    cases = [
        {"segment_tier": 1, "buyer_type": "WHOLESALER", "phone_normalized": "x", "email": "x@y", "website": "x", "social_links": ["x"], "governorate": "cairo", "last_verified_at": now, "source_last_seen_at": now},
        {"segment_tier": 1, "buyer_type": "RETAILER", "phone_normalized": "x", "email": "x@y", "website": "x", "social_links": [], "governorate": "cairo", "last_verified_at": None, "source_last_seen_at": now},
        {"segment_tier": 2, "buyer_type": "PROFESSIONAL_SERVICE", "phone_normalized": "x", "email": None, "website": None, "social_links": [], "governorate": "aswan", "last_verified_at": None, "source_last_seen_at": now - timedelta(days=60)},
        {"segment_tier": None, "buyer_type": None, "phone_normalized": None, "email": None, "website": None, "social_links": [], "governorate": "", "last_verified_at": None, "source_last_seen_at": now - timedelta(days=365)},
    ]
    assert [LeadScorer().score(case, source_count=1)[1] for case in cases] == ["A+", "A", "B", "C"]


def test_source_rate_limit_preserves_provider_evidence(monkeypatch) -> None:
    import httpx

    response = httpx.Response(
        429,
        text="request quota exceeded",
        headers={"Retry-After": "120"},
        request=httpx.Request("GET", "https://stac.overturemaps.org/catalog.json"),
    )
    monkeypatch.setattr(httpx, "get", lambda *args, **kwargs: response)
    with pytest.raises(SourceRateLimitError) as captured:
        OvertureCollector().discover_latest_release()
    assert captured.value.as_dict() == {
        "source": "SRC_OVERTURE",
        "http_code": 429,
        "provider_message": "request quota exceeded",
        "retry_after": "120",
        "reset_time": "RESET_TIME_UNKNOWN",
    }


def test_osm_bounded_overpass_sample_is_streamed(monkeypatch) -> None:
    import httpx

    payload = {
        "osm3s": {"timestamp_osm_base": "2026-08-25T00:00:00Z"},
        "elements": [{"type": "node", "id": 42, "lat": 30.04, "lon": 31.23, "tags": {"name": "Public Pharmacy", "amenity": "pharmacy", "phone": "+201000000000"}}],
    }
    response = httpx.Response(200, json=payload, request=httpx.Request("POST", "https://overpass-api.de/api/interpreter"))
    monkeypatch.setattr(httpx, "post", lambda *args, **kwargs: response)
    records = []
    context = CollectionContext(source_uid="SRC_OSM_GEOFABRIK", run_uid="r", job_uid="j", governorate="cairo", tile={"bbox": [31.2, 30.0, 31.3, 30.1]}, segment_ids=["pharmacies"], keywords=["pharmacy"], max_records=10, options={"osm_mode": "overpass"})
    result = OSMGeofabrikCollector().collect(context, lambda record, position: records.append(record) or True)
    assert result == {"release": "2026-08-25T00:00:00Z", "scanned": 1, "emitted": 1, "mode": "overpass_bbox"}
    assert records[0].source_record_id == "node/42"


def test_lead_api_workbook_pagination_controls_and_exports(api_request) -> None:
    dry = api_request("POST", "/lead-runs", json={"name": "api dry", "sources": ["SRC_OVERTURE"], "governorates": ["cairo"], "segments": ["pharmacies"], "dry_run": True, "execute": False})
    assert dry.status_code == 201 and dry.json()["planned_jobs"] > 0
    created = api_request("POST", "/lead-runs", json={"name": "api run", "sources": ["SRC_OVERTURE"], "governorates": ["cairo"], "segments": ["pharmacies"], "execute": False})
    assert created.status_code == 201
    uid = created.json()["run_uid"]
    assert api_request("POST", f"/lead-runs/{uid}/pause").status_code == 200
    assert api_request("POST", f"/lead-runs/{uid}/cancel").status_code == 200
    assert api_request("GET", "/lead-runs?offset=0&limit=10").json()["total"] == 1
    assert api_request("GET", "/leads?offset=0&limit=10").json()["total"] == 0
    for format_name in ("csv", "xlsx"):
        response = api_request("POST", "/leads/export", json={"format": format_name})
        assert response.status_code == 200 and response.content
    parquet = api_request("POST", "/leads/export", json={"format": "parquet"})
    parquet_status = api_request("GET", "/system/settings").json()["data_runtime"]["parquet"]
    if parquet_status["available"]:
        assert parquet.status_code == 200 and parquet.content
    else:
        assert parquet.status_code == 503
        assert parquet.json() == {"detail": "PARQUET_EXPORT_UNAVAILABLE"}

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("Lead Segments", "Keyword Master", "Egypt Coverage", "Source Registry", "Query Matrix", "Run Config"):
        sheet = workbook.create_sheet(name); sheet.append(["Keyword", "Approved"]); sheet.append(["pharmacy", True])
    payload = BytesIO(); workbook.save(payload)
    imported = api_request("POST", "/lead-control/workbook/import", files={"file": ("commands.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert imported.status_code == 201
    assert imported.json()["status"] == "VALID"
