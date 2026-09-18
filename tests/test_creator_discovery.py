from io import BytesIO
import time
from types import SimpleNamespace

from openpyxl import load_workbook

from backend.creator_discovery.analysis import EvidenceOnlyCreatorAnalyzer
from backend.creator_discovery.connectors import DirectURLConnector
from backend.creator_discovery.export import (
    REFERENCE_COLUMNS,
    build_creator_intelligence_xlsx,
    build_xlsx,
    export_rows,
)
from backend.creator_discovery.input import parse_creator_input
from backend.creator_discovery.matching import IdentityMatcher
from backend.creator_discovery.normalization import (
    detect_platform,
    normalize_creator_name,
    normalize_discovery_input,
    normalize_profile_url,
)
from backend.schemas.creator_discovery import DiscoveryPlatform
from backend.creator_discovery.workbook_schema import TEMPLATE_SHEETS


def _wait_for_run(api_request, run_uid: str, timeout: float = 60.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = api_request("GET", f"/creator-discovery/runs/{run_uid}")
        assert response.status_code == 200
        payload = response.json()
        if payload["status"] in {"COMPLETED", "PARTIAL", "FAILED", "CANCELLED"}:
            return payload
        time.sleep(0.05)
    raise AssertionError(f"Creator discovery run {run_uid} did not finish")


def test_name_url_and_platform_normalization():
    assert normalize_creator_name("  Ahmed—Example!  ") == "ahmed example"
    assert detect_platform("https://www.linkedin.com/in/ahmed-example") is DiscoveryPlatform.LINKEDIN
    platform, url, username = normalize_profile_url("http://twitter.com/Ahmed_Example/?ref=home")
    assert (platform, url, username) == (DiscoveryPlatform.X, "https://x.com/Ahmed_Example", "ahmed_example")
    direct = normalize_discovery_input("https://www.snapchat.com/add/Ahmed")
    assert direct.kind == "PROFILE_URL" and direct.username == "ahmed"


def test_identity_matching_does_not_confirm_on_name_alone():
    matcher = IdentityMatcher()
    name_only = matcher.match("Ahmed Example", {"display_name": "Ahmed Example"})
    assert name_only.confidence < 80
    complete = matcher.match("Ahmed Example", {"display_name": "Ahmed Example", "username": "ahmedexample"})
    assert complete.confidence >= 80
    cross_platform = matcher.match("Ahmed Example", {
        "display_name": "Ahmed Example", "username": "ahmedexample",
        "linked_social_urls": ["https://x.com/ahmed"],
        "reference_social_urls": ["https://x.com/ahmed"],
        "content_topics": ["technology", "reviews"],
        "reference_content_topics": ["technology", "reviews"],
        "location": "Cairo", "reference_location": "Cairo",
        "brand_names": ["EMY"], "reference_brand_names": ["EMY"],
    })
    assert cross_platform.confidence >= complete.confidence
    assert {item.signal for item in cross_platform.signals} >= {"linked_social_urls", "content_topic_similarity", "public_location", "brand_names", "profile_naming_consistency"}
    direct = matcher.match("https://www.youtube.com/@ahmed", {"profile_url": "https://www.youtube.com/@ahmed"}, direct_url=True)
    assert direct.confidence == 100 and direct.classification.value == "CONFIRMED"


def test_csv_and_xlsx_inputs_are_bounded_and_deduplicated():
    csv_values = parse_creator_input("names.csv", b"Name\nAhmed\nSara\nAhmed\n")
    assert csv_values == ["Ahmed", "Sara"]
    payload = build_xlsx([{"name": "Ahmed", "industry": "Education"}])
    assert parse_creator_input("names.xlsx", payload) == ["Ahmed"]
    mixed = b"Creator Name,Profile URL\nAhmed,https://www.instagram.com/ahmed\nSara,\n"
    assert parse_creator_input("mixed.csv", mixed) == ["https://www.instagram.com/ahmed", "Sara"]


def test_reference_export_mapping_and_summary():
    profile = {
        "name": "Ahmed", "username": "ahmed", "platform": "youtube",
        "niche": "Edutainment / Education", "industry": "Education",
        "main_platform": "youtube", "followers": 800000, "following": 120,
        "content_count": 85, "bio": "Educational technology", "verified": True,
        "profile_url": "https://www.youtube.com/@ahmed",
        "content_type": "VIDEO", "content_angle": "Education and problem solving",
        "content_style": "Educational explainers", "visual_style": "UNKNOWN",
        "cta": "subscribe", "engagement": "mean public engagement rate 5.00%",
        "audience_analysis": "public audience size 800,000; demographic data unavailable",
        "trend_analysis": "weekly", "ai_recommendation": "Prioritize explainers.",
        "personalization": "Tailor to learners.", "campaign_idea": "Tutorial series",
        "creative_strategy": "Repeatable explainers", "community_impact": "High sharing",
        "kpi_impact": "High educational authority", "source": "youtube_api", "confidence": 97.5,
        "last_updated": "2026-09-17T12:30:00+00:00", "x_url": "https://x.com/ahmed",
    }
    rows = export_rows([profile])
    assert list(rows[0]) == REFERENCE_COLUMNS and "X" not in rows[0]
    assert rows[0]["Creator Name"] == "Ahmed"
    assert rows[0]["Username"] == "ahmed"
    assert rows[0]["Profile URL"] == "https://www.youtube.com/@ahmed"
    workbook = load_workbook(BytesIO(build_xlsx([profile])))
    assert workbook.sheetnames == ["Creator Discovery", "Strategy Framework"]
    dump = workbook["Creator Discovery"]
    assert [cell.value for cell in dump[1]] == REFERENCE_COLUMNS
    assert dump.freeze_panes == "A2"
    assert dump.auto_filter.ref == "A1:AC2"
    assert dump["A1"].font.bold is True and dump["A1"].fill.fgColor.rgb == "001F4E78"
    assert dump.column_dimensions["A"].width > len("Creator Name")
    assert dump["AC2"].number_format == "yyyy-mm-dd hh:mm"
    assert dump["D2"].value == "https://www.youtube.com/@ahmed"
    assert dump["E2"].value == 800000
    assert workbook["Strategy Framework"][2][0].value == "User"
    assert export_rows([{**profile, "name": "=HYPERLINK(\"bad\")"}])[0]["Creator Name"].startswith("'=")

    extended = load_workbook(BytesIO(build_xlsx([profile], extended=True)))
    assert [cell.value for cell in extended["Creator Discovery"][1]] == [*REFERENCE_COLUMNS, "X"]

    intelligence = load_workbook(BytesIO(build_creator_intelligence_xlsx([{
        "profile": {**profile, "creator_uid": "CR_TEST", "analysis_status": "COMPLETED"},
        "accounts": [{
            "platform": "youtube", "username": "ahmed", "profile_url": profile["profile_url"],
            "followers": 800000, "first_seen_at": "2026-01-01T00:00:00", "metadata_json": {},
        }],
        "samples": [{
            "content_uid": "CDS_TEST", "platform": "youtube", "content_url": "https://youtube.com/watch?v=test",
            "published_at": "2026-08-01T12:00:00", "title": "Phone tutorial", "caption": "How to review a phone? Follow for more #technology",
            "content_type": "VIDEO", "views": 10000, "likes": 1000, "comments": 100, "shares": 50,
            "hashtags": ["technology"], "metadata_json": {"visual_style": "Studio demo"},
        }],
        "analysis": {
            **profile, "niche": profile["niche"], "cta_analysis": "subscribe",
            "engagement_pattern": profile["engagement"], "trend_analysis": profile["trend_analysis"],
            "community_impact": profile["community_impact"], "content_themes": ["technology"],
        },
        "sources": [{"source": "youtube_api", "platform": "youtube", "field_name": "followers"}],
    }])))
    assert intelligence.sheetnames == [name for name, _ in TEMPLATE_SHEETS]
    assert intelligence["01_Creator_Master_Database"]["A2"].value == "CR_TEST"
    assert intelligence["03_Content_Archive"]["A2"].value == "CDS_TEST"
    assert intelligence["31_Output_Creator_Master_Profile"]["C2"].value == "Ahmed"
    assert intelligence["58_VIDEO_TRANSCRIPT_INTELLIGENCE"]["A2"].value == "UNKNOWN"


def test_connector_fallback_and_evidence_only_industry_classification():
    connector = DirectURLConnector(DiscoveryPlatform.INSTAGRAM)
    outcome = connector.discover(normalize_discovery_input("https://www.instagram.com/ahmed"))
    assert outcome.status == "FOUND"
    assert outcome.candidates[0].discovery_source == "operator_supplied_url"
    assert outcome.candidates[0].followers is None

    account = SimpleNamespace(
        bio="Technology and gadget reviews", followers=None, content_count=None,
        metadata_json={}, platform="instagram", account_uid="CDA_TEST",
    )
    result = EvidenceOnlyCreatorAnalyzer().analyze(accounts=[account], samples=[])
    assert result["industry"] == "Technology"
    assert result["niche"] == "Technology / Gadgets"
    assert result["kpi_impact"] == "UNKNOWN"
    assert result["content_type"] == "UNKNOWN"
    assert result["audience_analysis"] == "UNKNOWN"
    assert result["ai_recommendation"] == "INSUFFICIENT_EVIDENCE"


def test_direct_url_run_review_merge_profile_and_export_api(api_request):
    created = api_request("POST", "/creator-discovery/runs", json={
        "inputs": ["https://www.youtube.com/@example"],
        "platforms": ["youtube", "instagram"],
        "analyze_content": True,
        "resolve_cross_platform_identity": True,
        "update_existing_profiles": False,
        "content_sample_size": 10,
        "auto_process": False,
        "execute": True,
    })
    assert created.status_code == 201
    completed_run = _wait_for_run(api_request, created.json()["run_uid"])
    assert completed_run["status"] == "PARTIAL"
    assert completed_run["matched"] == 1

    queue = api_request("GET", "/creator-discovery/candidates", params={"review_status": "PENDING"})
    assert queue.status_code == 200 and queue.json()["total"] == 1
    candidate = queue.json()["items"][0]
    assert candidate["platform"] == "youtube" and candidate["confidence"] == 100
    detail = api_request(
        "POST", f"/creator-discovery/candidates/{candidate['candidate_uid']}/confirm",
        json={"action": "THIS_IS_THE_ACCOUNT"},
    )
    assert detail.status_code == 200
    unified = detail.json()["unified_profile"]
    assert unified["youtube_url"] == "https://www.youtube.com/@example"
    assert unified["analysis_status"] == "INSUFFICIENT_EVIDENCE"
    initial_export = api_request("POST", "/creators/export", json={"format": "xlsx", "extended": False})
    assert initial_export.headers["content-disposition"] == 'attachment; filename="creator_discovery_export.xlsx"'
    initial_workbook = load_workbook(BytesIO(initial_export.content))
    initial_dump = initial_workbook["Creator Discovery"]
    assert initial_dump["A2"].value == unified["name"]
    assert initial_dump["B2"].value == "example"
    assert initial_dump["D2"].value == "https://www.youtube.com/@example"
    assert initial_dump["AA2"].value == "operator_supplied_url"
    corrected = api_request(
        "PATCH", f"/creator-discovery/creators/{unified['profile_uid']}",
        json={"niche": "Technology / Reviews", "main_platform": None, "start_year": None},
    )
    assert corrected.status_code == 200
    assert corrected.json()["unified_profile"]["niche"] == "Technology / Reviews"
    assert corrected.json()["unified_profile"]["main_platform"] == "UNKNOWN"

    second = api_request("POST", "/creator-discovery/runs", json={
        "inputs": ["https://www.instagram.com/example"], "platforms": ["instagram"],
        "auto_process": False,
        "execute": True,
    })
    assert second.status_code == 201
    assert _wait_for_run(api_request, second.json()["run_uid"])["status"] in {"COMPLETED", "PARTIAL"}
    second_queue = api_request("GET", "/creator-discovery/candidates", params={"review_status": "PENDING"}).json()
    instagram = next(item for item in second_queue["items"] if item["platform"] == "instagram")
    merged = api_request(
        "POST", f"/creator-discovery/candidates/{instagram['candidate_uid']}/confirm",
        json={"action": "MERGE", "creator_uid": unified["creator_uid"]},
    )
    assert merged.status_code == 200
    assert set(merged.json()["unified_profile"]["platforms_found"]) == {"youtube", "instagram"}

    result_page = api_request("GET", "/creator-discovery/creators", params={"platform": "instagram"})
    assert result_page.status_code == 200 and result_page.json()["total"] == 1
    exported = api_request("POST", "/creators/export", json={"format": "xlsx", "extended": False})
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook.sheetnames == ["Creator Discovery", "Strategy Framework"]


def test_bulk_run_auto_processes_valid_creators_and_exports_only_the_run(api_request):
    inputs = [
        f"https://www.tiktok.com/@bulk-{index}"
        for index in range(1, 21)
    ]
    created = api_request("POST", "/creator-discovery/runs", json={
        "inputs": inputs,
        "platforms": ["tiktok"],
        "mode": "bulk",
        "auto_process": True,
        "analyze_content": True,
        "resolve_cross_platform_identity": True,
        "execute": True,
    })
    assert created.status_code == 201
    run = _wait_for_run(api_request, created.json()["run_uid"])
    assert run["processed"] == 20
    assert run["matched"] == 20
    assert run["review_required"] == 0

    creators = api_request("GET", "/creator-discovery/creators")
    assert creators.status_code == 200
    assert creators.json()["total"] == 20

    exported = api_request("POST", "/creator-discovery/export", json={
        "format": "xlsx", "run_uid": run["run_uid"],
    })
    workbook = load_workbook(BytesIO(exported.content))
    dump = workbook["Creator Discovery"]
    assert dump.max_row == 21
    assert [cell.value for cell in dump[1]] == REFERENCE_COLUMNS
    assert dump["C2"].value == "tiktok"


def test_run_pause_resume_and_connector_capabilities(api_request):
    capabilities = api_request("GET", "/creator-discovery/connectors")
    assert capabilities.status_code == 200 and len(capabilities.json()) == 7
    assert all(item["name_discovery"] in {"API_REQUIRED", "MANUAL_URL_REQUIRED", "DIRECT_URL_REQUIRED", "CONFIGURED"} for item in capabilities.json())
    created = api_request("POST", "/creator-discovery/runs", json={
        "inputs": ["Ahmed Example"], "platforms": ["youtube"], "execute": False,
    }).json()
    paused = api_request("POST", f"/creator-discovery/runs/{created['run_uid']}/pause")
    assert paused.status_code == 200 and paused.json()["status"] == "PAUSED"
    resumed = api_request("POST", f"/creator-discovery/runs/{created['run_uid']}/resume")
    assert resumed.status_code == 200 and resumed.json()["status"] == "PENDING"
    assert _wait_for_run(api_request, created["run_uid"])["status"] == "PARTIAL"
