"""MSC and spreadsheet intake: file hashing, staging, review, and posting."""

from __future__ import annotations

import csv
import hashlib
import io
from abc import ABC, abstractmethod
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.core.config import get_settings
from backend.core.exceptions import ConflictError, NotFoundError, SalesValidationError
from backend.core.paths import paths
from backend.db.models.sales import MscIntakeBatch, MscIntakeFile, MscInvoiceStaging
from backend.repositories.sales_repository import SalesRepository
from backend.sales.common import stable_hash, uid
from backend.sales.invoicing import InvoiceService
from backend.sales.master_data import CustomerService, ProductService
from backend.schemas.sales import InvoiceCreate, InvoiceLineCreate, MscBatchCreate, MscExtractRequest, MscExtractionRecord, PostRequest


ALLOWED_MIME = {"image/png", "image/jpeg", "application/pdf", "text/csv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
ALLOWED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".pdf", ".csv", ".xlsx"}


class BaseMSCExtractor(ABC):
    code = "base"
    version = "1.0.0"

    @abstractmethod
    def extract_invoice(self, content: bytes, *, file_name: str) -> list[dict[str, Any]]: ...

    def extract_sales_report(self, content: bytes, *, file_name: str) -> list[dict[str, Any]]: return []
    def extract_stock_report(self, content: bytes, *, file_name: str) -> list[dict[str, Any]]: return []
    def extract_payment_report(self, content: bytes, *, file_name: str) -> list[dict[str, Any]]: return []
    def normalize_result(self, result: dict[str, Any]) -> dict[str, Any]: return dict(result)


class MockMSCExtractor(BaseMSCExtractor):
    """Deterministic no-network extractor used by tests and local validation."""
    code = "mock"

    def extract_invoice(self, content: bytes, *, file_name: str) -> list[dict[str, Any]]:
        return []


class MSCIntakeService:
    def __init__(self, session: Session, extractor: BaseMSCExtractor | None = None) -> None:
        self.repository = SalesRepository(session); self.extractor = extractor or MockMSCExtractor(); self.settings = get_settings()

    def create_batch(self, request: MscBatchCreate) -> MscIntakeBatch:
        key = request.idempotency_key or stable_hash(request.model_dump())
        existing = self.repository.msc_batch_by_key(key)
        if existing: return existing
        return self.repository.add(MscIntakeBatch(batch_uid=uid("MSC"), source_system="MSC", business_date=request.business_date, capture_date=request.capture_date, file_count=0, sequence_start=request.sequence_start, sequence_end=request.sequence_end, source_hash=request.source_hash.lower(), extractor=request.extractor, extractor_version=request.extractor_version, status="INTAKE", idempotency_key=key))

    def get(self, identifier: int | str) -> MscIntakeBatch:
        item = self.repository.get_msc_batch(identifier)
        if item is None: raise NotFoundError(f"MSC batch {identifier} was not found")
        return item

    def add_file(self, identifier: int | str, *, file_name: str, content: bytes, mime_type: str, sequence: int, capture_time: datetime | None = None, document_uid: str | None = None, page_number: int | None = None, page_count: int | None = None) -> MscIntakeFile:
        batch = self.get(identifier)
        suffix = Path(file_name).suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS or mime_type not in ALLOWED_MIME: raise SalesValidationError("Unsupported MSC file type")
        if len(content) > self.settings.import_max_file_size_mb * 1024 * 1024: raise SalesValidationError("MSC file exceeds configured size")
        digest = hashlib.sha256(content).hexdigest(); duplicate = self.repository.msc_file_by_hash(digest)
        if duplicate: raise ConflictError(f"Duplicate MSC file hash already exists as {duplicate.file_uid}")
        safe_name = f"{uid('MSCFILE')}{suffix}"
        base = paths.imports / "msc" / batch.batch_uid; base.mkdir(parents=True, exist_ok=True)
        target = paths.resolve_under(base, safe_name); target.write_bytes(content)
        relative = paths.validate_storage_value(paths.relative(target).as_posix())
        item = self.repository.add(MscIntakeFile(file_uid=uid("MSCF"), batch_id=batch.id, document_uid=document_uid, page_number=page_number, page_count=page_count, file_name=Path(file_name).name[:255], relative_path=relative, mime_type=mime_type, file_hash=digest, sequence=sequence, capture_time=capture_time, extraction_status="PENDING", validation_status="PENDING", duplicate_status="UNIQUE"))
        batch.file_count = len(self.repository.msc_files(batch.id)); batch.status = "FILES_RECEIVED"
        return item

    def stage_extraction(self, identifier: int | str, request: MscExtractRequest) -> list[MscInvoiceStaging]:
        batch = self.get(identifier); staged = []
        for record in request.records:
            source_file = self.repository.get_msc_file(record.source_file_id)
            if source_file is None or source_file.batch_id != batch.id: raise SalesValidationError("Extraction file does not belong to batch")
            raw = record.model_dump(); fingerprint = stable_hash({"batch": batch.batch_uid, **raw})
            existing = next((item for item in self.repository.msc_staging(batch.id) if item.fingerprint == fingerprint), None)
            if existing: staged.append(existing); continue
            fields = dict(record.fields); fields.update({"extractor": batch.extractor, "extractor_version": batch.extractor_version, "source_file_uid": source_file.file_uid})
            item = self.repository.add(MscInvoiceStaging(staging_uid=uid("STAGE"), batch_id=batch.id, source_file_id=source_file.id, invoice_number_raw=record.invoice_number, invoice_date_raw=record.invoice_date, customer_raw=record.customer, product_raw=record.product, qty_raw=record.quantity, price_raw=record.price, total_raw=record.total, extracted_fields=fields, confidence=record.confidence, validation_status="PENDING", fingerprint=fingerprint))
            source_file.extraction_status = "EXTRACTED"; staged.append(item)
        batch.status = "EXTRACTED"
        return staged

    def validate(self, identifier: int | str) -> dict:
        batch = self.get(identifier); files = self.repository.msc_files(batch.id); issues = []
        sequences = {item.sequence for item in files}
        if sequences:
            start = batch.sequence_start if batch.sequence_start is not None else min(sequences); end = batch.sequence_end if batch.sequence_end is not None else max(sequences)
            for value in range(start, end + 1):
                if value not in sequences: issues.append({"code": f"MISSING_SEQUENCE_{value}", "severity": "BLOCK"})
        for item in self.repository.msc_staging(batch.id):
            product = self.repository.get_product(item.product_raw.strip()) if item.product_raw else None
            if product is None and item.product_raw:
                matches = [candidate for candidate in self.repository.list_products() if item.product_raw.casefold() in {(candidate.product_name_en or "").casefold(), (candidate.product_name_ar or "").casefold()}]
                product = matches[0] if len(matches) == 1 else None
            customer = self.repository.get_customer(item.customer_raw.strip()) if item.customer_raw else None
            if customer is None and item.customer_raw:
                matches = [candidate for candidate in self.repository.list_customers() if item.customer_raw.casefold() in {candidate.business_name.casefold(), (candidate.trade_name or "").casefold()}]
                customer = matches[0] if len(matches) == 1 else None
            item.product_id = product.id if product else None; item.customer_id = customer.id if customer else None
            item.product_match = "EXACT" if product else "NO_MATCH"; item.customer_match = "EXACT" if customer else "NO_MATCH"
            row_issues = []
            if item.confidence < Decimal(str(self.settings.msc_high_confidence_threshold)): row_issues.append("LOW_CONFIDENCE")
            if product is None: row_issues.append("UNKNOWN_PRODUCT")
            if customer is None: row_issues.append("AMBIGUOUS_OR_UNKNOWN_CUSTOMER")
            try:
                qty = Decimal(item.qty_raw or ""); price = Decimal(item.price_raw or ""); total = Decimal(item.total_raw or "")
                if qty <= 0 or price < 0 or total < 0: raise InvalidOperation
                if abs((qty * price) - total) > Decimal("0.01"): row_issues.append("TOTAL_MISMATCH")
                date.fromisoformat(item.invoice_date_raw or "")
            except (InvalidOperation, ValueError): row_issues.append("INVALID_ARITHMETIC_OR_DATE")
            if not item.invoice_number_raw: row_issues.append("MISSING_INVOICE_NUMBER")
            item.validation_status = "VALID" if not row_issues else "HUMAN_REVIEW"
            issues.extend({"staging_uid": item.staging_uid, "code": code, "severity": "BLOCK"} for code in row_issues)
        batch.status = "VALIDATED" if not issues else "REVIEW_REQUIRED"
        for file in files: file.validation_status = "VALID" if not issues else "REVIEW_REQUIRED"
        return {"batch": batch, "issues": issues, "valid": sum(item.validation_status == "VALID" for item in self.repository.msc_staging(batch.id)), "invalid": sum(item.validation_status != "VALID" for item in self.repository.msc_staging(batch.id))}

    def approve(self, identifier: int | str, approved_by: str) -> MscIntakeBatch:
        batch = self.get(identifier)
        validation = self.validate(batch.id)
        if validation["issues"]: raise SalesValidationError("MSC batch requires human correction before approval")
        for item in self.repository.msc_staging(batch.id): item.validation_status = "APPROVED"; item.approved_by = approved_by
        batch.status = "APPROVED"
        return batch

    def post(self, identifier: int | str, posted_by: str) -> MscIntakeBatch:
        batch = self.get(identifier)
        if batch.status == "POSTED": return batch
        if batch.status != "APPROVED": raise SalesValidationError("No posting from unapproved staging")
        invoice_service = InvoiceService(self.repository.session)
        for item in self.repository.msc_staging(batch.id):
            if item.posted_invoice_id: continue
            request = InvoiceCreate(invoice_number=item.invoice_number_raw, customer_id=item.customer_id, invoice_date=date.fromisoformat(item.invoice_date_raw), currency=str(item.extracted_fields.get("currency", "EGP")), source_system="MSC", idempotency_key=f"msc:{batch.batch_uid}:{item.fingerprint}", lines=[InvoiceLineCreate(product_id=item.product_id, warehouse_id=item.extracted_fields.get("warehouse_id"), quantity=Decimal(item.qty_raw), unit_price=Decimal(item.price_raw))])
            if request.lines[0].warehouse_id is None: raise SalesValidationError("MSC posting requires an approved warehouse match")
            invoice = invoice_service.create(request); invoice_service.post(invoice.id, PostRequest(posted_by=posted_by)); item.posted_invoice_id = invoice.id; item.validation_status = "POSTED"
        batch.status = "POSTED"
        return batch


class SalesSheetImportService:
    REQUIRED = {"invoice_number", "invoice_date", "customer", "sku", "quantity", "unit_price", "warehouse"}

    @staticmethod
    def _rows(content: bytes, filename: str) -> list[dict[str, Any]]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv": return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        if suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True); sheet = workbook.active
            rows = sheet.iter_rows(values_only=True); headers = [str(value or "").strip() for value in next(rows)]
            return [dict(zip(headers, values, strict=False)) for values in rows]
        raise SalesValidationError("Sales Sheet must be CSV or XLSX")

    def preview(self, content: bytes, filename: str) -> dict:
        rows = self._rows(content, filename); valid = []; invalid = []
        for number, row in enumerate(rows, start=2):
            missing = sorted(field for field in self.REQUIRED if row.get(field) in {None, ""})
            target = valid if not missing else invalid; target.append({"row": number, "data": row, "errors": [f"MISSING_{field.upper()}" for field in missing]})
        return {"dry_run": True, "records": len(rows), "valid": len(valid), "invalid": len(invalid), "duplicates": 0, "unknown_customers": [], "unknown_products": [], "warnings": [], "valid_rows": valid, "invalid_rows": invalid}
