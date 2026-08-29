"""Excel command-workbook discovery, validation, and managed import."""

from __future__ import annotations

import hashlib
from pathlib import Path
import re
from typing import Any, ClassVar
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.core.paths import paths
from backend.db.models.leads import LeadControlImport
from backend.leads.config import LeadCatalog, get_lead_catalog


class LeadWorkbookService:
    _preview_cache: ClassVar[dict[tuple[str, int], dict[str, Any]]] = {}
    _rows_cache: ClassVar[dict[tuple[str, int, str], list[dict[str, Any]]]] = {}
    HEADER_MARKERS = {
        "Lead Segments": "category_id",
        "Keyword Master": "keyword_id",
        "Egypt Coverage": "gov_id",
        "Source Registry": "source_id",
        "Query Matrix": "query_id",
        "Run Config": "config_id",
    }

    def __init__(self, session: Session | None = None, catalog: LeadCatalog | None = None) -> None:
        self.session = session
        self.catalog = catalog or get_lead_catalog()

    @property
    def placement(self) -> str:
        return (Path(self.catalog.workbook.get("directory", "data/imports/lead_control")) / "EMY_Egypt_Lead_Acquisition_OS_v1.xlsx").as_posix()

    def find(self) -> Path | None:
        patterns = self.catalog.workbook.get("patterns", ["EMY_Egypt_Lead_Acquisition_OS_v*.xlsx"])
        for directory in (paths.lead_control_imports, paths.project_root):
            for pattern in patterns:
                matches = sorted(directory.glob(str(pattern)), key=lambda item: item.stat().st_mtime, reverse=True)
                if matches:
                    return matches[0]
        return None

    def preview(self, workbook_path: Path | None = None) -> dict[str, Any]:
        candidate = workbook_path or self.find()
        required = list(self.catalog.workbook.get("required_sheets", []))
        if candidate is None:
            return {
                "available": False,
                "status": "NOT_FOUND_USING_CONFIG_DEFAULTS",
                "active_command_source": "CONFIG_DEFAULTS",
                "placement": self.placement,
                "required_sheets": required,
                "sheet_counts": {},
                "query_matrix_rows": 0,
                "enabled_query_jobs": 0,
                "errors": [],
            }
        preview_key = (str(candidate.resolve()), candidate.stat().st_mtime_ns)
        if preview_key in self._preview_cache:
            return self._preview_cache[preview_key]
        workbook = load_workbook(candidate, read_only=True, data_only=True)
        resolved = {name: self._resolve_sheet_name(workbook.sheetnames, name) for name in required}
        missing = [name for name, actual in resolved.items() if actual is None]
        sheet_rows = {}
        for name, actual in resolved.items():
            if actual is None:
                continue
            parsed = self._rows_from_sheet(workbook[actual], name)
            sheet_rows[name] = parsed
            self._rows_cache[(str(candidate.resolve()), candidate.stat().st_mtime_ns, name)] = parsed
        counts = {name: len(rows) for name, rows in sheet_rows.items()}
        query_rows = sheet_rows.get("Query Matrix", [])
        enabled_query_jobs = sum(1 for row in query_rows if self._enabled(row.get("Enabled", True)))
        workbook.close()
        try:
            stored_path = paths.relative(candidate).as_posix()
        except ValueError:
            stored_path = candidate.name
        result = {
            "available": True,
            "status": "VALID" if not missing else "INVALID",
            "active_command_source": "WORKBOOK" if not missing else "CONFIG_DEFAULTS",
            "filename": candidate.name,
            "stored_path": stored_path,
            "placement": self.placement,
            "required_sheets": required,
            "resolved_sheets": {name: actual for name, actual in resolved.items() if actual is not None},
            "sheet_counts": counts,
            "query_matrix_rows": len(query_rows),
            "enabled_query_jobs": enabled_query_jobs,
            "errors": [f"Missing sheet: {name}" for name in missing],
        }
        self._preview_cache[preview_key] = result
        return result

    def rows(self, sheet_name: str, workbook_path: Path | None = None) -> list[dict[str, Any]]:
        candidate = workbook_path or self.find()
        if candidate is None:
            return []
        cache_key = (str(candidate.resolve()), candidate.stat().st_mtime_ns, sheet_name)
        if cache_key in self._rows_cache:
            return self._rows_cache[cache_key]
        workbook = load_workbook(candidate, read_only=True, data_only=True)
        actual_name = self._resolve_sheet_name(workbook.sheetnames, sheet_name)
        if actual_name is None:
            workbook.close()
            return []
        result = self._rows_from_sheet(workbook[actual_name], sheet_name)
        workbook.close()
        self._rows_cache[cache_key] = result
        return result

    def effective_keywords(self) -> list[dict[str, Any]]:
        if self.preview()["status"] != "VALID":
            return self.catalog.approved_keywords
        rows = self.rows("Keyword Master")
        category_ids = self._segment_id_map()
        parsed = []
        for row in rows:
            normalized = self._normalized_row(row)
            keyword = normalized.get("keyword") or normalized.get("term") or normalized.get("search_phrase")
            approved = normalized.get("approved", True)
            status = normalized.get("status", "ENABLED")
            raw_category = str(normalized.get("category_id") or normalized.get("segment_id") or "")
            if keyword and self._enabled(approved) and self._enabled(status):
                parsed.append({
                    "keyword": str(keyword),
                    "category_id": category_ids.get(raw_category, raw_category or None),
                    "approved": True,
                    "language": normalized.get("language") or normalized.get("lang"),
                })
        return parsed

    def effective_segments(self) -> list[dict[str, Any]]:
        if self.preview()["status"] != "VALID":
            return self.catalog.segments
        keyword_terms: dict[str, list[str]] = {}
        for row in self.rows("Keyword Master"):
            normalized = self._normalized_row(row)
            category_id = str(normalized.get("category_id") or "")
            keyword = normalized.get("keyword")
            if category_id and keyword and self._enabled(normalized.get("status", "ENABLED")):
                keyword_terms.setdefault(category_id, []).append(str(keyword))
        result = []
        for row in self.rows("Lead Segments"):
            normalized = self._normalized_row(row)
            workbook_id = str(normalized.get("category_id") or "")
            name = str(normalized.get("segment_en") or normalized.get("segment") or workbook_id)
            try:
                phase = int(normalized.get("phase") or 1)
            except (TypeError, ValueError):
                phase = 1
            result.append({
                "category_id": self._runtime_segment_id(name, workbook_id),
                "workbook_category_id": workbook_id,
                "name": name,
                "name_ar": str(normalized.get("segment_ar") or ""),
                "tier": phase,
                "fit_class": str(normalized.get("tier") or ""),
                "buyer_type": str(normalized.get("buyer_type") or ""),
                "terms": keyword_terms.get(workbook_id, []),
            })
        return result

    def effective_governorates(self) -> list[dict[str, Any]]:
        if self.preview()["status"] != "VALID":
            from backend.leads.geography import GeographyService

            return GeographyService(self.catalog).list_governorates()
        configured = {self._key(item["name"]): item for item in self.catalog.governorates}
        name_aliases = {"fayoum": "faiyum", "assiut": "asyut"}
        result = []
        for row in self.rows("Egypt Coverage"):
            normalized = self._normalized_row(row)
            name = str(normalized.get("governorate_en") or "")
            name_key = self._key(name)
            base = configured.get(name_key) or configured.get(name_aliases.get(name_key, ""))
            if base is None:
                continue
            from backend.leads.geography import GeographyService

            bbox = tuple(float(value) for value in base["bbox"])
            result.append({
                **base,
                "name": name,
                "name_ar": str(normalized.get("governorate_ar") or base.get("name_ar") or ""),
                "workbook_governorate_id": str(normalized.get("gov_id") or ""),
                "polygon": [list(point) for point in GeographyService._bbox_polygon(bbox)],
            })
        return result

    def _segment_id_map(self) -> dict[str, str]:
        result = {}
        for row in self.rows("Lead Segments"):
            normalized = self._normalized_row(row)
            workbook_id = str(normalized.get("category_id") or "")
            name = str(normalized.get("segment_en") or normalized.get("segment") or workbook_id)
            result[workbook_id] = self._runtime_segment_id(name, workbook_id)
        return result

    def _runtime_segment_id(self, name: str, fallback: str) -> str:
        match = next((item for item in self.catalog.segments if self._key(item["name"]) == self._key(name)), None)
        return str(match["category_id"]) if match else self._slug(name) or fallback

    def _resolve_sheet_name(self, sheet_names: list[str], requested: str) -> str | None:
        if requested in sheet_names:
            return requested
        aliases = self.catalog.workbook.get("sheet_aliases", {})
        alias = aliases.get(requested)
        if alias in sheet_names:
            return str(alias)
        requested_key = self._key(requested)
        return next((name for name in sheet_names if self._key(re.sub(r"^\d+[_\s-]*", "", name)) == requested_key), None)

    def _rows_from_sheet(self, sheet: Any, canonical_name: str) -> list[dict[str, Any]]:
        materialized = list(sheet.iter_rows(values_only=True))
        marker = self.HEADER_MARKERS.get(canonical_name)
        header_index = None
        for index, values in enumerate(materialized):
            normalized = {self._key(value) for value in values if value not in (None, "")}
            if marker and marker in normalized:
                header_index = index
                break
        if header_index is None:
            header_index = next((index for index, values in enumerate(materialized) if any(value not in (None, "") for value in values)), None)
        if header_index is None:
            return []
        headers = [str(value).strip() if value is not None else "" for value in materialized[header_index]]
        result = []
        for values in materialized[header_index + 1:]:
            item = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]}
            if any(value not in (None, "") for value in item.values()):
                result.append(item)
        return result

    @classmethod
    def _normalized_row(cls, row: dict[str, Any]) -> dict[str, Any]:
        return {cls._key(key): value for key, value in row.items()}

    @staticmethod
    def _key(value: Any) -> str:
        return re.sub(r"[^a-z0-9]+", "_", str(value).strip().casefold()).strip("_")

    @classmethod
    def _slug(cls, value: str) -> str:
        return cls._key(value)

    @staticmethod
    def _enabled(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        return str(value).strip().casefold() not in {"", "0", "false", "no", "n", "disabled", "inactive"}

    def import_bytes(self, filename: str, content: bytes) -> dict[str, Any]:
        if Path(filename).suffix.casefold() != ".xlsx":
            raise ValueError("Lead command workbook must be an .xlsx file")
        if len(content) > 25 * 1024 * 1024:
            raise ValueError("Lead command workbook exceeds the 25 MB limit")
        paths.lead_control_imports.mkdir(parents=True, exist_ok=True)
        safe_name = f"EMY_Egypt_Lead_Acquisition_OS_v1_{uuid4().hex[:10]}.xlsx"
        destination = paths.lead_control_imports / safe_name
        destination.write_bytes(content)
        preview = self.preview(destination)
        if self.session is not None:
            record = LeadControlImport(
                import_uid=f"LCI_{uuid4().hex.upper()}",
                filename=Path(filename).name,
                stored_path=paths.relative(destination).as_posix(),
                sha256=hashlib.sha256(content).hexdigest(),
                status=preview["status"],
                sheet_counts=preview["sheet_counts"],
                errors=preview["errors"],
            )
            self.session.add(record)
            self.session.flush()
            preview["import_uid"] = record.import_uid
        return preview

