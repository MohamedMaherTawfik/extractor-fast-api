"""Bounded CSV/XLSX creator input parsing."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook


INPUT_HEADERS = (
    "Profile URL", "URL", "Account URL",
    "Creator Name", "Name", "Account Name", "Creator", "Account",
)


def parse_creator_input(filename: str, content: bytes, *, limit: int = 10_000) -> list[str]:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        return _parse_csv(content, limit)
    if suffix == ".xlsx":
        return _parse_xlsx(content, limit)
    raise ValueError("Creator input must be a CSV or XLSX file")


def _parse_csv(content: bytes, limit: int) -> list[str]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV input has no header row")
    headers = [name for name in INPUT_HEADERS if name in reader.fieldnames]
    if not headers:
        headers = [reader.fieldnames[0]]
    return _bounded_unique((_first_value(row.get(header) for header in headers) for row in reader), limit)


def _parse_xlsx(content: bytes, limit: int) -> list[str]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook["EGYPT_MASSIVE_DUMP"] if "EGYPT_MASSIVE_DUMP" in workbook.sheetnames else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if not headers:
            raise ValueError("XLSX input has no header row")
        indexes = [headers.index(name) for name in INPUT_HEADERS if name in headers]
        if not indexes:
            indexes = [0]
        return _bounded_unique((
            _first_value(str(row[index]) if len(row) > index and row[index] is not None else "" for index in indexes)
            for row in rows
        ), limit)
    finally:
        workbook.close()


def _bounded_unique(values, limit: int) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = value.strip()
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        result.append(cleaned)
        if len(result) > limit:
            raise ValueError(f"Creator input exceeds the {limit} row limit")
    if not result:
        raise ValueError("Creator input contains no names or profile URLs")
    return result


def _first_value(values) -> str:
    return next((str(value).strip() for value in values if value is not None and str(value).strip()), "")

