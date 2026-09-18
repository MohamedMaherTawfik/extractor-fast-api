"""Exact workbook-compatible Creator Discovery exports."""

from __future__ import annotations

import csv
import json
import statistics
from collections import Counter
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.creator_discovery.workbook_schema import TEMPLATE_SHEETS


REFERENCE_COLUMNS = [
    "Creator Name", "Username", "Platform", "Profile URL", "Followers", "Following",
    "Content Count", "Bio", "Verified", "Industry", "Niche", "Main Platform",
    "Content Type", "Content Angle", "Content Style", "Visual Style", "CTA",
    "Engagement", "Audience Analysis", "Trend Analysis", "AI Recommendation",
    "Personalization", "Campaign Idea", "Creative Strategy", "Community Impact",
    "KPI Impact", "Source", "Confidence", "Last Updated",
]
EXTENDED_COLUMNS = [*REFERENCE_COLUMNS, "X"]
CSV_COLUMNS = [
    "Name", "Niche", "Industry", "Main_Platform", "YouTube", "Facebook",
    "Instagram", "TikTok", "Snapchat", "LinkedIn", "Content_Mechanism_Style",
    "Influence_Size", "KPI_Impact", "Start_Year",
]
EXTENDED_CSV_COLUMNS = [*CSV_COLUMNS, "X"]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))
BODY_FONT = Font(name="Arial", size=10, color="1F2937")


def export_rows(profiles: list[dict[str, Any]], *, extended: bool = False) -> list[dict[str, Any]]:
    rows = []
    for item in profiles:
        row = {
            "Creator Name": item.get("name"),
            "Username": item.get("username"),
            "Platform": item.get("platform"),
            "Profile URL": item.get("profile_url"),
            "Followers": item.get("followers"),
            "Following": item.get("following"),
            "Content Count": item.get("content_count"),
            "Bio": item.get("bio"),
            "Verified": item.get("verified"),
            "Industry": item.get("industry"),
            "Niche": item.get("niche"),
            "Main Platform": item.get("main_platform"),
            "Content Type": item.get("content_type"),
            "Content Angle": item.get("content_angle"),
            "Content Style": item.get("content_style"),
            "Visual Style": item.get("visual_style"),
            "CTA": item.get("cta"),
            "Engagement": item.get("engagement"),
            "Audience Analysis": item.get("audience_analysis"),
            "Trend Analysis": item.get("trend_analysis"),
            "AI Recommendation": item.get("ai_recommendation"),
            "Personalization": item.get("personalization"),
            "Campaign Idea": item.get("campaign_idea"),
            "Creative Strategy": item.get("creative_strategy"),
            "Community Impact": item.get("community_impact"),
            "KPI Impact": item.get("kpi_impact"),
            "Source": item.get("source"),
            "Confidence": item.get("confidence"),
            "Last Updated": _excel_datetime(item.get("last_updated") or item.get("updated_at")),
        }
        if extended:
            row["X"] = item.get("x_url")
        rows.append({column: _spreadsheet_safe(value) for column, value in row.items()})
    return rows


def build_xlsx(profiles: list[dict[str, Any]], *, extended: bool = False) -> bytes:
    rows = export_rows(profiles, extended=extended)
    columns = EXTENDED_COLUMNS if extended else REFERENCE_COLUMNS
    workbook = Workbook()
    dump = workbook.active
    dump.title = "Creator Discovery"
    dump.append(columns)
    for row in rows:
        dump.append([row[column] for column in columns])
    _format_worksheet(dump, columns)

    framework = workbook.create_sheet("Strategy Framework")
    framework_rows = [
        ["Category", "Element", "Description", "Data Source", "AI Analysis Required"],
        ["User", "Creator Profile", "Profile identity and creator information", None, "Yes"],
        ["Content", "Content Type", "Type and format of published content", None, "Yes"],
        ["Content", "Content Angle", "Creative direction and messaging angle", None, "Yes"],
        ["AI", "Recommendation", "AI-based suggestions", None, "Yes"],
        ["Audience", "Audience Analysis", "Audience behavior and characteristics", None, "Yes"],
        ["Marketing", "CTA", "Call to action analysis", None, "Yes"],
        ["Marketing", "KPI", "Performance measurement", None, "Yes"],
        ["Platform", "Social Platform", "Platform presence and performance", None, "Yes"],
        ["Strategy", "Campaign Idea", "Campaign concepts generated from insights", None, "Yes"],
        ["Strategy", "Creative Strategy", "Content planning approach", None, "Yes"],
    ]
    for row in framework_rows:
        framework.append(row)
    _format_worksheet(framework, framework_rows[0], freeze=False)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_creator_intelligence_xlsx(records: list[dict[str, Any]]) -> bytes:
    """Build the complete multi-sheet workbook without changing the legacy export."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows_by_sheet = _intelligence_rows(records)
    for sheet_name, columns in TEMPLATE_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(list(columns))
        rows = rows_by_sheet.get(sheet_name) or [["UNKNOWN" for _ in columns]]
        for row in rows:
            padded = [*row[:len(columns)], *(["UNKNOWN"] * max(0, len(columns) - len(row)))]
            worksheet.append([_spreadsheet_safe(_excel_datetime(value)) for value in padded])
        _format_worksheet(worksheet, list(columns))
        _format_intelligence_numbers(worksheet, list(columns))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _intelligence_rows(records: list[dict[str, Any]]) -> dict[str, list[list[Any]]]:
    result: dict[str, list[list[Any]]] = {name: [] for name, _ in TEMPLATE_SHEETS}
    for record in records:
        profile = record["profile"]
        accounts = record.get("accounts") or []
        samples = record.get("samples") or []
        intelligence = record.get("analysis") or {}
        sources = record.get("sources") or []
        creator_uid = profile.get("creator_uid") or "UNKNOWN"
        creator_name = profile.get("name") or "UNKNOWN"
        usernames = sorted({_text(item.get("username")) for item in accounts if item.get("username")})
        urls = {str(item.get("platform")): item.get("profile_url") for item in accounts}
        dated_samples = [item for item in samples if item.get("published_at")]
        earliest = (
            min(dated_samples, key=lambda item: _sort_time(item.get("published_at"))).get("published_at")
            if dated_samples else None
        )
        first_platform = min(accounts, key=lambda item: _sort_time(item.get("first_seen_at")))["platform"] if accounts else "UNKNOWN"
        content_themes = _themes(samples, intelligence)
        hooks = [_hook(item.get("caption") or item.get("title")) for item in samples]
        cta = _text(intelligence.get("cta_analysis"))
        content_style = _text(intelligence.get("content_style") or profile.get("content_mechanism_style"))
        visual_style = _text(intelligence.get("visual_style"))
        audience = _text(intelligence.get("audience_analysis"))
        recommendation = _text(intelligence.get("ai_recommendation"))
        campaign = _text(intelligence.get("campaign_idea"))
        strategy = _text(intelligence.get("creative_strategy"))
        engagement = _text(intelligence.get("engagement_pattern"))

        result["01_Creator_Master_Database"].append([
            creator_uid, "UNKNOWN", creator_name, ", ".join(usernames) or "UNKNOWN",
            _metadata(accounts, "country"), _metadata(accounts, "city"), _metadata(accounts, "language"),
            profile.get("start_year") or "UNKNOWN", first_platform, _sample_url(min(samples, key=lambda x: _sort_time(x.get("published_at")))) if samples else "UNKNOWN",
            profile.get("analysis_status") or "UNKNOWN", urls.get("youtube") or "UNKNOWN", urls.get("facebook") or "UNKNOWN",
            urls.get("instagram") or "UNKNOWN", urls.get("tiktok") or "UNKNOWN", urls.get("snapchat") or "UNKNOWN",
            urls.get("linkedin") or "UNKNOWN", "UNKNOWN", "UNKNOWN",
        ])
        result["02_Creator_History_2014_NOW"].append([
            profile.get("start_year") or "UNKNOWN", "First observed creator evidence", first_platform,
            content_style, _text(intelligence.get("growth_signals")), "UNKNOWN", _text(profile.get("industry")),
        ])
        for sample in samples:
            rate = _engagement_rate(sample, accounts)
            topic = _sample_topic(sample, intelligence)
            row = [
                sample.get("content_uid"), creator_name, sample.get("platform"), sample.get("published_at"),
                sample.get("title") or _hook(sample.get("caption")), sample.get("content_url"), sample.get("content_type"),
                topic, sample.get("views"), sample.get("likes"), sample.get("comments"), sample.get("shares"),
                sample.get("metadata_json", {}).get("saves"), _performance_label(rate),
            ]
            result["03_Content_Archive"].append(row)
            result["35_Output_Content_Inventory"].append([
                sample.get("content_uid"), sample.get("title") or _hook(sample.get("caption")), sample.get("content_url"),
                sample.get("platform"), sample.get("published_at"), sample.get("duration_seconds"), sample.get("content_type"),
                topic, sample.get("views"), sample.get("likes"), sample.get("comments"), sample.get("shares"),
                sample.get("metadata_json", {}).get("saves"), rate,
            ])
            caption = sample.get("caption") or "UNKNOWN"
            hook = _hook(caption)
            metadata = sample.get("metadata_json") or {}
            result["04_Video_Forensic_Analysis"].append([
                sample.get("content_url"), hook, _text(intelligence.get("storytelling_style")),
                _text(intelligence.get("content_angle")), _text(intelligence.get("retention_mechanism")),
                metadata.get("camera") or "UNKNOWN", metadata.get("lens") or "UNKNOWN", metadata.get("lighting") or "UNKNOWN",
                metadata.get("editing_style") or visual_style, metadata.get("sound") or "UNKNOWN", cta,
                _performance_reason(sample, rate, intelligence),
            ])
            result["11_Hook_Database"].append([
                hook, _hook_type(hook), _text(intelligence.get("audience_trigger")), profile.get("industry"),
                sample.get("content_url"), _performance_label(rate),
            ])
            result["42_Output_Hook_Analysis"].append([
                hook, metadata.get("visual_hook") or "UNKNOWN", metadata.get("spoken_hook") or "UNKNOWN",
                _hook_type(hook), "YES" if "?" in hook else "UNKNOWN", _text(intelligence.get("audience_problem")),
                _text(intelligence.get("content_angle")), "UNKNOWN", _text(intelligence.get("storytelling_style")),
                metadata.get("pattern_interrupt") or "UNKNOWN",
            ])
            result["44_Output_Video_Deconstruction"].append([
                hook, "UNKNOWN", caption, _closing_copy(caption), metadata.get("shot_type") or "UNKNOWN",
                metadata.get("camera_angle") or "UNKNOWN", metadata.get("composition") or "UNKNOWN",
                metadata.get("lens") or "UNKNOWN", metadata.get("movement") or "UNKNOWN", metadata.get("cuts") or "UNKNOWN",
                metadata.get("transitions") or "UNKNOWN", caption, metadata.get("music") or "UNKNOWN", metadata.get("sound_effects") or "UNKNOWN",
            ])
            result["51_Video_Copywriting_Archive"].append([
                sample.get("content_uid"), creator_name, sample.get("platform"), sample.get("content_url"), sample.get("published_at"),
                sample.get("title") or "UNKNOWN", caption, metadata.get("description") or "UNKNOWN",
                ", ".join(sample.get("hashtags") or []) or "UNKNOWN", hook, caption, _closing_copy(caption), cta,
                _text(intelligence.get("caption_style")), _text(intelligence.get("copy_framework")),
                _text(intelligence.get("communication_style")), _text(intelligence.get("persuasion_elements")),
            ])
            result["52_Video_By_Video_Full_Analysis"].append([
                sample.get("content_uid"), hook, metadata.get("visual_hook") or "UNKNOWN", metadata.get("spoken_hook") or "UNKNOWN",
                _text(intelligence.get("storytelling_style")), caption, metadata.get("transcript") or "UNKNOWN",
                metadata.get("scene_count") or "UNKNOWN", metadata.get("scenes") or "UNKNOWN", metadata.get("camera_direction") or "UNKNOWN",
                metadata.get("editing_style") or visual_style, metadata.get("transitions") or "UNKNOWN", metadata.get("music") or "UNKNOWN",
                metadata.get("sound_effects") or "UNKNOWN", visual_style, _text(intelligence.get("audience_trigger")),
                _performance_reason(sample, rate, intelligence),
            ])
            result["53_Video_Copy_Reverse_Engineering"].append([
                sample.get("content_uid"), topic, _text(intelligence.get("content_angle")), audience,
                _text(intelligence.get("audience_problem")), _text(intelligence.get("value_promise")), caption,
                _text(intelligence.get("emotional_trigger")), _text(intelligence.get("persuasion_elements")), cta,
                _text(intelligence.get("creative_pattern")),
            ])

        result["06_Creator_DNA"].append([
            _text(intelligence.get("personality")), _text(intelligence.get("voice")),
            _text(intelligence.get("tone")), _text(intelligence.get("presentation_style")),
            _text(intelligence.get("storytelling_style")), visual_style, _text(intelligence.get("editing_style")),
            _text(intelligence.get("audience_relationship")), _text(intelligence.get("success_formula")),
        ])
        result["07_Audience_Intelligence"].append([
            _text(intelligence.get("audience_type")), "UNKNOWN", _metadata(accounts, "audience_location"),
            _text(intelligence.get("audience_needs")), _text(intelligence.get("audience_problem")),
            _text(intelligence.get("audience_motivation")), _text(intelligence.get("engagement_behavior")),
            _text(intelligence.get("audience_questions")), _text(intelligence.get("audience_trigger")),
        ])
        result["08_Strategy_Engine"].append([
            _text(intelligence.get("mission")), _text(intelligence.get("vision")),
            _text(intelligence.get("positioning")), _text(intelligence.get("usp")),
            _text(intelligence.get("benchmark")), _text(intelligence.get("differentiation")),
        ])
        for theme, count in content_themes:
            result["09_Content_Pillars"].append([
                theme, count / max(sum(value for _, value in content_themes), 1),
                _text(intelligence.get("campaign_goal")), _theme_examples(theme, samples),
            ])
        result["10_Idea_Generation_Engine"].append([
            campaign, _text(intelligence.get("audience_problem")), _text(intelligence.get("content_angle")),
            _text(intelligence.get("audience_trigger")), hooks[0] if hooks else "UNKNOWN",
            profile.get("main_platform") or "UNKNOWN", _text(intelligence.get("content_type")),
        ])
        for account in accounts:
            platform_samples = [item for item in samples if item.get("platform") == account.get("platform")]
            views = [item.get("views") for item in platform_samples if item.get("views") is not None]
            ranked = sorted(platform_samples, key=lambda item: _engagement_score(item), reverse=True)
            result["18_Platform_Algorithm"].append([
                account.get("platform"), _text(intelligence.get("platform_factors")), "UNKNOWN", "UNKNOWN",
                sum(views) if views else "UNKNOWN", "UNKNOWN", "UNKNOWN",
                sum(item.get("shares") or 0 for item in platform_samples) if platform_samples else "UNKNOWN",
                sum((item.get("metadata_json") or {}).get("saves") or 0 for item in platform_samples) if platform_samples else "UNKNOWN",
            ])
            result["34_Output_Platform_Audit"].append([
                account.get("platform"), account.get("profile_url"), _metadata([account], "creation_date"), account.get("followers"),
                account.get("followers") if account.get("platform") == "youtube" else "UNKNOWN",
                sum(views) if views else "UNKNOWN", round(statistics.fmean(views), 2) if views else "UNKNOWN",
                _sample_url(ranked[0]) if ranked else "UNKNOWN", _sample_url(ranked[-1]) if ranked else "UNKNOWN",
                _posting_frequency(platform_samples), ", ".join(theme for theme, _ in content_themes[:3]) or "UNKNOWN",
            ])
            result["62_SOCIAL_MEDIA_INTELLIGENCE_ENGINE"].append([
                account.get("platform"), "Profile, content, engagement, and positioning",
                engagement, content_style, recommendation,
            ])
            result["71_PLATFORM_OPTIMIZATION_COMMANDS"].append([
                account.get("platform"), _text(intelligence.get("posting_behavior")),
                _text(intelligence.get("platform_factors")), recommendation,
            ])
        total_views = sum(item.get("views") or 0 for item in samples) if any(item.get("views") is not None for item in samples) else "UNKNOWN"
        result["19_Analytics_KPI"].append([
            max((item.get("followers") or 0 for item in accounts), default="UNKNOWN"), total_views,
            engagement, _text(intelligence.get("growth_signals")), "UNKNOWN", "UNKNOWN", "UNKNOWN",
        ])
        result["22_Trend_Event_Intelligence"].append([
            _text(intelligence.get("trend_analysis")), earliest or "UNKNOWN", profile.get("main_platform"),
            "Collected content samples", _text(intelligence.get("trend_impact")), _text(intelligence.get("growth_opportunities")),
        ])
        result["24_Community_Management"].append([
            _text(intelligence.get("community_impact")), "UNKNOWN", _text(intelligence.get("sentiment")),
            "UNKNOWN", _text(intelligence.get("community_signals")),
        ])
        for source in sources:
            result["27_Data_Source_Collection"].append([
                source.get("source"), source.get("platform") or "unified", "Official API or recorded provenance",
                source.get("field_name"),
            ])
        result["30_Final_Creator_Intelligence_Report"].extend([
            ["Creator positioning", "Profile bio and sampled captions", _text(intelligence.get("positioning")), recommendation],
            ["Content pattern", f"{len(samples)} stored content samples", content_style, strategy],
            ["Audience", audience, _text(intelligence.get("audience_type")), _text(intelligence.get("brand_suitability"))],
            ["Growth", _text(intelligence.get("trend_analysis")), _text(intelligence.get("growth_signals")), _text(intelligence.get("growth_opportunities"))],
        ])
        result["31_Output_Creator_Master_Profile"].append([
            creator_uid, "UNKNOWN", creator_name, ", ".join(usernames) or "UNKNOWN", "CREATOR",
            _metadata(accounts, "country"), _metadata(accounts, "city"), _metadata(accounts, "language"), "UNKNOWN",
            profile.get("industry"), profile.get("niche"),
        ])
        result["32_Output_Origin_Story"].append([
            profile.get("start_year") or "UNKNOWN", earliest or "UNKNOWN", first_platform,
            _sample_url(min(samples, key=lambda x: _sort_time(x.get("published_at")))) if samples else "UNKNOWN",
            _sample_topic(samples[-1], intelligence) if samples else "UNKNOWN", "UNKNOWN", "UNKNOWN", "UNKNOWN",
        ])
        result["33_Output_Complete_Timeline"].append([
            profile.get("start_year") or "UNKNOWN", "First observed evidence", first_platform,
            _text(intelligence.get("content_type")), max((item.get("followers") or 0 for item in accounts), default="UNKNOWN"),
            _text(intelligence.get("growth_signals")), len(samples), "UNKNOWN", "UNKNOWN", "UNKNOWN", strategy,
        ])
        ranked_samples = sorted(samples, key=_engagement_score, reverse=True)
        for index, sample in enumerate(ranked_samples, start=1):
            result["36_Output_Top_Performing_Content"].append([
                index, sample.get("content_url"), sample.get("views"), sample.get("shares"), sample.get("comments"),
                (sample.get("metadata_json") or {}).get("saves"), "UNKNOWN",
                _performance_reason(sample, _engagement_rate(sample, accounts), intelligence), _hook(sample.get("caption")),
                _text(intelligence.get("storytelling_style")), _text(intelligence.get("emotional_trigger")),
                _text(intelligence.get("trend_analysis")),
            ])
        classifications = _classifications(samples)
        result["37_Output_Content_Classification"].append([
            classifications.get(key, 0) for key in (
                "educational", "entertainment", "storytelling", "review", "reaction",
                "documentary", "vlog", "podcast", "tutorial", "challenge",
            )
        ])
        result["38_Output_Content_DNA"].append([
            _text(intelligence.get("personality")), _text(intelligence.get("voice")), _text(intelligence.get("energy")),
            _text(intelligence.get("presentation_style")), _text(intelligence.get("editing_style")), visual_style,
            _text(intelligence.get("audience_relationship")), _text(intelligence.get("success_formula")),
        ])
        result["39_Output_Audience_Analysis"].append([
            "UNKNOWN", "UNKNOWN", _metadata(accounts, "audience_location"), _text(intelligence.get("audience_interests")),
            _text(intelligence.get("audience_problem")), _text(intelligence.get("audience_needs")),
            _text(intelligence.get("engagement_behavior")), "UNKNOWN", _text(intelligence.get("audience_type")),
        ])
        result["40_Output_Buyer_Persona"].append([
            _text(intelligence.get("audience_type")), "UNKNOWN", "UNKNOWN", _text(intelligence.get("audience_problem")),
            "UNKNOWN", _text(intelligence.get("audience_needs")), _text(intelligence.get("audience_motivation")),
            "UNKNOWN", _text(intelligence.get("audience_relationship")),
        ])
        result["41_Output_Content_Strategy"].append([
            _text(intelligence.get("positioning")), _text(intelligence.get("usp")), _text(intelligence.get("mission")),
            _text(intelligence.get("vision")), _text(intelligence.get("brand_identity")),
            _text(intelligence.get("differentiation")),
        ])
        result["48_Output_UGC_Analysis"].append([
            visual_style, _text(intelligence.get("personality")), _text(intelligence.get("brand_suitability")),
            _text(intelligence.get("audience_problem")), _text(intelligence.get("value_promise")),
            engagement, _text(intelligence.get("community_signals")), cta,
        ])
        result["49_Output_Copywriting"].append([
            _text(intelligence.get("aida")), _text(intelligence.get("pas")), _text(intelligence.get("fab")),
            _text(intelligence.get("four_u")), hooks[0] if hooks else "UNKNOWN",
            samples[0].get("caption") if samples else "UNKNOWN", _text(intelligence.get("value_promise")), cta,
        ])
        result["50_Output_CTA_Database"].append([
            cta, _text(intelligence.get("cta_type")), _text(intelligence.get("campaign_goal")),
            profile.get("main_platform"), engagement, "UNKNOWN", "UNKNOWN",
            "YES" if "follow" in cta.casefold() or "subscribe" in cta.casefold() else "UNKNOWN",
            "YES" if "subscribe" in cta.casefold() else "UNKNOWN",
        ])
        result["57_CREATOR_COMPLETE_COMMAND_PIPELINE"].extend([
            [1, "Discovery", "Resolve creator identity", "Creator URL or name", "Official connector", creator_uid],
            [2, "Collection", "Collect profile and content", "Platform account", "API normalization", f"{len(samples)} samples"],
            [3, "Analysis", "Build evidence-based intelligence", "Profile and samples", "Deterministic analysis", profile.get("analysis_status")],
            [4, "Strategy", "Generate recommendations", "Analysis output", "Evidence mapping", recommendation],
            [5, "Report", "Build workbook", "Persisted creator intelligence", "Template mapping", "Creator Intelligence workbook"],
        ])
        result["59_CONTENT_MACHINE_MASTER_OUTPUT"].append([
            creator_name, profile.get("start_year") or "UNKNOWN", ", ".join(urls) or "UNKNOWN",
            content_style, len(samples), "UNKNOWN", _text(intelligence.get("caption_style")),
            _text(intelligence.get("storytelling_style")), visual_style, "UNKNOWN", engagement,
            _text(intelligence.get("creative_pattern")),
        ])
        result["63_USER_EXPERIENCE_CONTENT_ENGINE"].append([
            _text(intelligence.get("audience_relationship")), _text(intelligence.get("audience_journey")),
            _text(intelligence.get("audience_problem")), engagement, recommendation,
        ])
        result["64_EGC_UGC_WOM_ENGINE"].append([
            _text(intelligence.get("content_type")), creator_name, _text(intelligence.get("community_signals")),
            _text(intelligence.get("word_of_mouth")), _text(intelligence.get("community_impact")),
        ])
        result["65_CONTENT_RECOMMENDATION_AI_ENGINE"].append([
            f"{len(samples)} content samples and {len(accounts)} platform accounts", "Content strategy",
            strategy, _text(intelligence.get("recommendation_reasoning")), recommendation,
        ])
        result["66_WORLD_OF_MARKETING_ENGINE"].append([
            "Creator collaboration", _text(intelligence.get("copy_framework")),
            _text(intelligence.get("collaboration_opportunities")), campaign,
        ])
        result["67_COPYWRITING_DEEP_ANALYSIS"].append([
            _text(intelligence.get("caption_style")), hooks[0] if hooks else "UNKNOWN", hooks[0] if hooks else "UNKNOWN",
            _text(intelligence.get("value_promise")), _text(intelligence.get("emotional_trigger")),
            _text(intelligence.get("persuasion_elements")), cta,
        ])
        result["68_HIDDEN_MESSAGES_ANALYSIS"].append([
            content_style, "UNKNOWN", "UNKNOWN", audience, _text(intelligence.get("brand_identity")),
        ])
        result["69_CREATIVE_STRATEGY_ENGINE"].append([
            campaign, _text(intelligence.get("content_angle")), strategy, visual_style,
            _text(profile.get("kpi_impact")),
        ])
        result["70_VIDEO_QUALITY_ENGINE"].append([
            _text(intelligence.get("content_type")), "UNKNOWN", visual_style,
            _text(intelligence.get("editing_style")), "Only API metadata and captions were available",
        ])
        result["26_Output_Report_Structure"].append([
            creator_name, "Creator, content, audience, and strategy intelligence", f"{len(TEMPLATE_SHEETS)} sheets",
            profile.get("analysis_status"),
        ])
    return result


def build_csv(profiles: list[dict[str, Any]], *, extended: bool = False) -> bytes:
    rows = _export_csv_rows(profiles, extended=extended)
    columns = EXTENDED_CSV_COLUMNS if extended else CSV_COLUMNS
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def _export_csv_rows(profiles: list[dict[str, Any]], *, extended: bool = False) -> list[dict[str, Any]]:
    rows = []
    for item in profiles:
        row = {
            "Name": item.get("name"),
            "Niche": item.get("niche") or "UNKNOWN",
            "Industry": item.get("industry") or "UNKNOWN",
            "Main_Platform": item.get("main_platform") or "UNKNOWN",
            "YouTube": item.get("youtube_url") or "NOT_AVAILABLE",
            "Facebook": item.get("facebook_url") or "NOT_AVAILABLE",
            "Instagram": item.get("instagram_url") or "NOT_AVAILABLE",
            "TikTok": item.get("tiktok_url") or "NOT_AVAILABLE",
            "Snapchat": item.get("snapchat_url") or "NOT_AVAILABLE",
            "LinkedIn": item.get("linkedin_url") or "NOT_AVAILABLE",
            "Content_Mechanism_Style": item.get("content_mechanism_style") or "UNKNOWN",
            "Influence_Size": item.get("influence_size") or "UNKNOWN",
            "KPI_Impact": item.get("kpi_impact") or "UNKNOWN",
            "Start_Year": item.get("start_year") or "UNKNOWN",
        }
        if extended:
            row["X"] = item.get("x_url") or "NOT_AVAILABLE"
        rows.append({column: _spreadsheet_safe(value) for column, value in row.items()})
    return rows


def _text(value: Any) -> str:
    if value is None:
        return "UNKNOWN"
    if isinstance(value, list):
        shown = ", ".join(str(item) for item in value if item not in (None, ""))
        return shown or "UNKNOWN"
    shown = str(value).strip()
    return shown if shown and shown not in {"[]", "{}"} else "UNKNOWN"


def _metadata(accounts: list[dict[str, Any]], key: str) -> Any:
    for account in accounts:
        value = (account.get("metadata_json") or {}).get(key)
        if value not in (None, "", [], {}):
            return value
    return "UNKNOWN"


def _sample_url(sample: dict[str, Any]) -> str:
    return _text(sample.get("content_url"))


def _sort_time(value: Any) -> datetime:
    converted = _excel_datetime(value)
    return converted if isinstance(converted, datetime) else datetime.max


def _engagement_score(sample: dict[str, Any]) -> int:
    return sum(int(sample.get(field) or 0) for field in ("views", "likes", "comments", "shares"))


def _engagement_rate(sample: dict[str, Any], accounts: list[dict[str, Any]]) -> float | str:
    followers = max((int(item.get("followers") or 0) for item in accounts), default=0)
    denominator = int(sample.get("views") or 0) or followers
    if not denominator:
        return "UNKNOWN"
    interactions = sum(int(sample.get(field) or 0) for field in ("likes", "comments", "shares"))
    return interactions / denominator


def _performance_label(rate: float | str) -> str:
    if not isinstance(rate, (int, float)):
        return "UNKNOWN"
    if rate >= 0.05:
        return "HIGH"
    if rate >= 0.02:
        return "MEDIUM"
    return "LOW"


def _hook(value: Any) -> str:
    text = _text(value)
    if text == "UNKNOWN":
        return text
    first_line = text.splitlines()[0].strip()
    sentence = first_line.split(".", 1)[0].strip()
    return (sentence or first_line)[:180]


def _closing_copy(value: Any) -> str:
    text = _text(value)
    if text == "UNKNOWN":
        return text
    sentences = [item.strip() for item in text.replace("\n", " ").split(".") if item.strip()]
    return sentences[-1][:240] if sentences else text[-240:]


def _hook_type(hook: str) -> str:
    lowered = hook.casefold()
    if "?" in hook:
        return "Question"
    if any(term in lowered for term in ("how to", "tips", "guide", "learn")):
        return "Educational"
    if any(term in lowered for term in ("why", "secret", "mistake", "never")):
        return "Curiosity"
    return "Statement" if hook != "UNKNOWN" else "UNKNOWN"


def _themes(samples: list[dict[str, Any]], intelligence: dict[str, Any]) -> list[tuple[str, int]]:
    hashtags = Counter(
        str(tag).strip().casefold()
        for sample in samples for tag in (sample.get("hashtags") or [])
        if str(tag).strip()
    )
    if hashtags:
        return hashtags.most_common(8)
    themes = intelligence.get("content_themes")
    if isinstance(themes, list) and themes:
        return [(str(item), 1) for item in themes[:8]]
    niche = _text(intelligence.get("niche"))
    return [(niche, 1)] if niche != "UNKNOWN" else [("UNKNOWN", 1)]


def _sample_topic(sample: dict[str, Any], intelligence: dict[str, Any]) -> str:
    tags = sample.get("hashtags") or []
    if tags:
        return ", ".join(str(item) for item in tags[:5])
    return _text(intelligence.get("niche"))


def _theme_examples(theme: str, samples: list[dict[str, Any]]) -> str:
    matches = [
        _sample_url(sample) for sample in samples
        if theme.casefold() in " ".join(str(tag).casefold() for tag in sample.get("hashtags") or [])
    ]
    return ", ".join(matches[:3]) or "UNKNOWN"


def _posting_frequency(samples: list[dict[str, Any]]) -> str:
    dates = sorted(
        _excel_datetime(item.get("published_at")) for item in samples if item.get("published_at")
    )
    dates = [item for item in dates if isinstance(item, datetime)]
    if len(dates) < 2:
        return "UNKNOWN"
    gaps = [(right - left).total_seconds() / 86400 for left, right in zip(dates, dates[1:])]
    return f"Median {statistics.median(gaps):.1f} days"


def _performance_reason(
    sample: dict[str, Any], rate: float | str, intelligence: dict[str, Any],
) -> str:
    facts = []
    if isinstance(rate, (int, float)):
        facts.append(f"observed engagement rate {rate:.2%}")
    if sample.get("views") is not None:
        facts.append(f"{int(sample['views']):,} views")
    trend = _text(intelligence.get("trend_analysis"))
    if trend != "UNKNOWN":
        facts.append(trend)
    return "; ".join(facts) or "UNKNOWN"


def _classifications(samples: list[dict[str, Any]]) -> Counter[str]:
    groups = {
        "educational": ("learn", "explained", "education", "tips", "how to"),
        "entertainment": ("funny", "comedy", "entertainment"),
        "storytelling": ("story", "journey", "experience"),
        "review": ("review", "unboxing", "comparison"),
        "reaction": ("reaction", "react"),
        "documentary": ("documentary",),
        "vlog": ("vlog",),
        "podcast": ("podcast",),
        "tutorial": ("tutorial", "how to"),
        "challenge": ("challenge",),
    }
    counts: Counter[str] = Counter()
    for sample in samples:
        text = f"{sample.get('title') or ''} {sample.get('caption') or ''}".casefold()
        for label, terms in groups.items():
            counts[label] += int(any(term in text for term in terms))
    return counts


def _format_intelligence_numbers(worksheet, columns: list[str]) -> None:
    percent_columns = {"Percentage", "Engagement_Rate", "CTR", "Retention", "Completion", "Rewatch", "ROI", "ROAS"}
    count_columns = {
        "Followers", "Subscribers", "Total_Views", "Average_Views", "Views", "Likes", "Comments",
        "Shares", "Saves", "Reach", "Audience_Size", "Videos", "Scene_Count",
    }
    date_columns = {"Date", "Publish_Date", "Creation_Date", "Start_Date", "First_Appearance", "Timestamp"}
    for index, name in enumerate(columns, start=1):
        cells = worksheet[get_column_letter(index)][1:]
        if name in percent_columns:
            for cell in cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0%"
        elif name in count_columns:
            for cell in cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
        elif name in date_columns:
            for cell in cells:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"


def _format_worksheet(worksheet, columns: list[str], *, freeze: bool = True) -> None:
    last_column = get_column_letter(len(columns))
    worksheet.freeze_panes = "A2" if freeze else None
    worksheet.auto_filter.ref = f"A1:{last_column}{max(worksheet.max_row, 1)}"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")

    for index, column_name in enumerate(columns, start=1):
        values = [worksheet.cell(row=row, column=index).value for row in range(1, worksheet.max_row + 1)]
        width = max((_display_length(value) for value in values), default=len(column_name)) + 2
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width, 12), 60)

    for count_column in ("Followers", "Following", "Content Count"):
        if count_column in columns:
            for cell in worksheet[get_column_letter(columns.index(count_column) + 1)][1:]:
                cell.number_format = "#,##0"
    if "Confidence" in columns:
        for cell in worksheet[get_column_letter(columns.index("Confidence") + 1)][1:]:
            cell.number_format = "0.0"
    if "Last Updated" in columns:
        for cell in worksheet[get_column_letter(columns.index("Last Updated") + 1)][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"


def _display_length(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime):
        return 16
    return max((len(line) for line in str(value).splitlines()), default=0)


def _excel_datetime(value: Any) -> Any:
    if value is None or isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value
    else:
        return value
    if parsed is not None and parsed.tzinfo is not None:
        return parsed.astimezone(UTC).replace(tzinfo=None)
    return parsed


def _spreadsheet_safe(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str, sort_keys=True)
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value
