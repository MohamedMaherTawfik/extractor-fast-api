"""Master-control authoring import, preview, classification, and safe compilation."""

from __future__ import annotations

from abc import ABC, abstractmethod
from datetime import UTC, datetime
from hashlib import sha256
import json
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.core.enums import (
    ControlClassification, RuleAction, RuleHardness, RuleKind, RuleLifecycleStatus,
    RuleSeverity, RuleSourceType, RuleStage,
)
from backend.core.exceptions import RuleValidationError
from backend.core.paths import paths
from backend.repositories.rule_repository import RuleRepository
from backend.schemas.rules import (
    CompilableControl, MasterControlPreviewResponse, RuleActionSpec, RuleCreate,
)
from backend.services.rule_registry_service import RuleRegistryService, stable_hash


EXECUTABLE_CLASSIFICATIONS = {
    ControlClassification.CONSTRAINT, ControlClassification.VALIDATION_RULE,
    ControlClassification.POLICY, ControlClassification.QA_RULE,
    ControlClassification.ACCESSIBILITY_RULE, ControlClassification.PLATFORM_RULE,
    ControlClassification.BRAND_RULE, ControlClassification.CHARACTER_RULE,
}


class MasterControlImporter(ABC):
    @abstractmethod
    def read(self, source: Path, sheet_name: str | None = None) -> tuple[str, list[dict[str, Any]]]:
        """Return the selected sheet and normalized row mappings."""


class ExcelMasterControlImporter(MasterControlImporter):
    def read(self, source: Path, sheet_name: str | None = None):
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            selected = sheet_name or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise RuleValidationError(f"Master control sheet not found: {selected}")
            sheet = workbook[selected]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return selected, []
            keys = [self._key(value) for value in headers]
            return selected, [
                {keys[index]: value for index, value in enumerate(row) if index < len(keys) and keys[index] and value is not None}
                for row in rows
            ]
        finally:
            workbook.close()

    @staticmethod
    def _key(value):
        return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


class RuleCompiler:
    """Compile only explicit, complete executable controls."""

    TYPE_BY_CLASSIFICATION = {
        ControlClassification.CONSTRAINT: RuleKind.HARD_CONSTRAINT,
        ControlClassification.VALIDATION_RULE: RuleKind.VALIDATION,
        ControlClassification.POLICY: RuleKind.POLICY,
        ControlClassification.QA_RULE: RuleKind.QA_GATE,
        ControlClassification.ACCESSIBILITY_RULE: RuleKind.ACCESSIBILITY,
        ControlClassification.PLATFORM_RULE: RuleKind.PLATFORM,
        ControlClassification.BRAND_RULE: RuleKind.BRAND,
        ControlClassification.CHARACTER_RULE: RuleKind.CHARACTER,
    }

    def compile(self, control: CompilableControl) -> RuleCreate | None:
        if control.status.upper() in {"NEEDS_DEFINITION", "DISABLED_UNDEFINED"} or control.classification is ControlClassification.UNDEFINED:
            return None
        if control.classification is ControlClassification.INTEGRATION_ONLY:
            return None
        if control.classification not in EXECUTABLE_CLASSIFICATIONS or not control.condition or not control.action or not control.rule_code:
            return None
        try:
            action = RuleActionSpec.model_validate(control.action)
        except Exception as exc:
            raise RuleValidationError(f"Control {control.control_id} has an invalid controlled action") from exc
        return RuleCreate(
            rule_code=control.rule_code, name=control.name,
            description=f"Compiled from master control {control.control_id}",
            domain=control.domain or control.classification.value,
            rule_type=self.TYPE_BY_CLASSIFICATION[control.classification],
            hardness=control.hardness, severity=control.severity, priority=control.priority,
            condition=control.condition, action=action, message=control.name,
            source_control_id=control.control_id, source_type=RuleSourceType.MASTER_SHEET,
            source_reference=control.source_reference, status=RuleLifecycleStatus.DRAFT,
            stage=RuleStage.PRE_GENERATION,
        )


class MasterControlImportService:
    def __init__(self, session: Session, importer: MasterControlImporter | None = None) -> None:
        self.session = session
        self.repository = RuleRepository(session)
        self.registry = RuleRegistryService(session)
        self.importer = importer or ExcelMasterControlImporter()
        self.compiler = RuleCompiler()

    def preview(self, relative_path: str, *, sheet_name: str | None = None) -> MasterControlPreviewResponse:
        source = self._source(relative_path)
        if source is None or not source.is_file():
            return MasterControlPreviewResponse(source_file_hash=None, status="MASTER_CONTROL_FILE_NOT_FOUND")
        source_hash = self._hash_file(source)
        selected, rows = self.importer.read(source, sheet_name)
        prior_import = self.repository.find_import(source_hash, selected)
        changes = [self._classify_row(row, index + 2) for index, row in enumerate(rows) if row]
        if prior_import:
            for change in changes:
                change["change"] = "UNCHANGED"
        return self._summary(source_hash, "IDEMPOTENT" if prior_import else "PREVIEW", changes)

    def import_controls(self, relative_path: str, *, sheet_name: str | None = None, dry_run: bool = True):
        preview = self.preview(relative_path, sheet_name=sheet_name)
        if dry_run or preview.status in {"MASTER_CONTROL_FILE_NOT_FOUND", "IDEMPOTENT"}:
            return preview
        source = self._source(relative_path)
        selected, rows = self.importer.read(source, sheet_name)
        changes = [self._classify_row(row, index + 2) for index, row in enumerate(rows) if row]
        run = self.repository.create_import(
            import_uid=f"RIMPORT_{uuid4().hex}", source_file_hash=preview.source_file_hash,
            source_path=source.relative_to(paths.project_root).as_posix(), sheet_name=selected,
            status="completed", summary=preview.model_dump(mode="json"), imported_at=datetime.now(UTC),
        )
        for change in changes:
            row = change["payload"]
            latest = self.repository.latest_control(change["control_id"])
            if latest and latest.fingerprint == change["fingerprint"]:
                continue
            self.repository.create_control(
                control_id=change["control_id"], version=(latest.version + 1 if latest else 1),
                fingerprint=change["fingerprint"], classification=change["classification"],
                status=change["status"], source_row=change["source_row"], payload=change["payload"],
                import_run_id=run.id,
            )
            control = self._to_control(change)
            compiled = self.compiler.compile(control)
            existing_rule = self.repository.get(str(row.get("rule_code") or "")) if row.get("rule_code") else None
            if compiled and existing_rule is None:
                self.registry.create_rule(compiled)
            elif compiled and existing_rule is not None:
                from backend.schemas.rules import RuleVersionCreate
                self.registry.add_version(
                    existing_rule.id,
                    RuleVersionCreate.model_validate(compiled.model_dump(exclude={"rule_code", "status"})),
                )
            elif existing_rule is not None:
                for rule_version in existing_rule.versions:
                    if rule_version.status is RuleLifecycleStatus.ACTIVE:
                        self.registry.retire(existing_rule.id, rule_version.version)
        self.session.flush()
        return self._summary(preview.source_file_hash, "IMPORTED", changes)

    def _classify_row(self, row, source_row):
        control_id = str(row.get("control_id") or row.get("id") or f"ROW_{source_row}").strip()
        raw_classification = str(row.get("classification") or row.get("control_type") or "undefined").strip().lower()
        try:
            classification = ControlClassification(raw_classification)
        except ValueError:
            classification = ControlClassification.UNDEFINED
        status = str(row.get("status") or "active").strip().upper()
        if status in {"NEEDS_DEFINITION", "يحتاج_تعريف"}:
            classification = ControlClassification.UNDEFINED
            status = "DISABLED_UNDEFINED"
        executable = classification in EXECUTABLE_CLASSIFICATIONS and bool(row.get("condition")) and bool(row.get("action"))
        if classification is ControlClassification.INTEGRATION_ONLY:
            status = "INTEGRATION_ONLY"
        elif classification is ControlClassification.UNDEFINED:
            status = "DISABLED_UNDEFINED"
        elif not executable:
            status = "NON_EXECUTABLE_CONTROL"
        payload = {key: self._json_value(value) for key, value in row.items()}
        fingerprint = stable_hash(payload)
        latest = self.repository.latest_control(control_id)
        change = "NEW" if latest is None else ("UNCHANGED" if latest.fingerprint == fingerprint else "MODIFIED")
        return {
            "control_id": control_id, "classification": classification.value, "status": status,
            "source_row": source_row, "fingerprint": fingerprint, "change": change,
            "executable": executable and status == "ACTIVE", "payload": payload,
        }

    def _to_control(self, change):
        row = change["payload"]
        def enum(enum_type, value, default):
            try: return enum_type(str(value).lower())
            except ValueError: return default
        return CompilableControl(
            control_id=change["control_id"], name=str(row.get("name") or change["control_id"]),
            classification=ControlClassification(change["classification"]), status=change["status"],
            rule_code=row.get("rule_code"), domain=row.get("domain"),
            condition=self._mapping(row.get("condition")), action=self._mapping(row.get("action")),
            priority=int(row.get("priority") or 0),
            severity=enum(RuleSeverity, row.get("severity"), RuleSeverity.MEDIUM),
            hardness=enum(RuleHardness, row.get("hardness"), RuleHardness.SOFT),
            source_reference=f"{change['control_id']} row {change['source_row']}",
        )

    @staticmethod
    def _mapping(value):
        if isinstance(value, dict): return value
        if isinstance(value, str):
            try:
                parsed = json.loads(value)
                return parsed if isinstance(parsed, dict) else None
            except json.JSONDecodeError:
                return None
        return None

    @staticmethod
    def _json_value(value):
        if isinstance(value, (str, int, float, bool)) or value is None: return value
        if isinstance(value, datetime): return value.isoformat()
        return str(value)

    @staticmethod
    def _hash_file(source):
        digest = sha256()
        with source.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _source(relative_path):
        try:
            return paths.resolve_under(paths.project_root, relative_path)
        except ValueError:
            raise RuleValidationError("Master control path must stay inside the project")

    @staticmethod
    def _summary(source_hash, status, changes):
        count = lambda value: sum(item["change"] == value for item in changes)
        return MasterControlPreviewResponse(
            source_file_hash=source_hash, status=status, new_controls=count("NEW"),
            modified_controls=count("MODIFIED"), unchanged_controls=count("UNCHANGED"),
            disabled_controls=sum(item["status"] in {"DISABLED_UNDEFINED", "INTEGRATION_ONLY"} for item in changes),
            undefined_controls=sum(item["status"] == "DISABLED_UNDEFINED" for item in changes),
            executable_candidates=sum(item["executable"] for item in changes),
            non_executable_controls=sum(not item["executable"] for item in changes), changes=changes,
        )
