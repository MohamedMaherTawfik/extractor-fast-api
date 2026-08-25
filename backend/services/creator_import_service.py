"""Config-driven CSV/XLSX Creator Master importer."""

import csv
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import StringIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.core.config import get_settings
from backend.core.enums import ImportStatus, Platform
from backend.core.exceptions import ConflictError, ImportValidationError
from backend.core.paths import paths
from backend.db.models.creator import Creator
from backend.db.models.import_batch import CreatorImportBatch
from backend.repositories.creator_import_repository import CreatorImportRepository
from backend.repositories.creator_repository import CreatorRepository
from backend.repositories.platform_account_repository import (
    PlatformAccountRepository,
)
from backend.schemas.creator import CreatorCreate
from backend.services.normalization import AccountNormalizer, NormalizedAccount


CREATOR_FIELDS = {
    "creator_uid",
    "display_name",
    "country",
    "primary_language",
    "category",
    "notes",
    "priority",
    "active",
}


def normalize_header(value: str) -> str:
    return re.sub(r"[\W_]+", "", value.strip().casefold(), flags=re.UNICODE)


class ColumnMapping:
    def __init__(self, overrides: dict[str, str] | None = None) -> None:
        configured = dict(get_settings().creator_import_column_mapping)
        if overrides:
            configured.update(overrides)
        self._mapping = {
            normalize_header(source): target.strip()
            for source, target in configured.items()
            if source.strip() and target.strip()
        }

    def map_row(self, row: dict[str, Any]) -> dict[str, Any]:
        mapped: dict[str, Any] = {}
        for source, value in row.items():
            target = self._mapping.get(normalize_header(str(source)))
            if target and not self._is_empty(value):
                if target not in mapped or self._is_empty(mapped[target]):
                    mapped[target] = value
        return mapped

    @staticmethod
    def _is_empty(value: Any) -> bool:
        return value is None or str(value).strip() == ""


@dataclass(slots=True)
class ParsedImportRow:
    creator_values: dict[str, Any]
    accounts: list[NormalizedAccount]


class CreatorImportService:
    def __init__(self, session: Session) -> None:
        self._session = session
        self._imports = CreatorImportRepository(session)
        self._creators = CreatorRepository(session)
        self._accounts = PlatformAccountRepository(session)
        self._normalizer = AccountNormalizer()

    def import_bytes(
        self,
        *,
        filename: str,
        content: bytes,
        mapping_overrides: dict[str, str] | None = None,
    ) -> CreatorImportBatch:
        safe_name = Path(filename).name
        suffix = Path(safe_name).suffix.casefold()
        if suffix not in {".csv", ".xlsx"}:
            raise ImportValidationError("Only .csv and .xlsx files are supported")
        max_bytes = get_settings().import_max_file_size_mb * 1024 * 1024
        if not content or len(content) > max_bytes:
            raise ImportValidationError("Import file is empty or exceeds the size limit")

        batch_uid = f"IMP_{datetime.now(UTC):%Y%m%d}_{uuid4().hex[:12].upper()}"
        paths.ensure_runtime_directories()
        destination = paths.resolve_under(
            paths.imports,
            f"{batch_uid}_{safe_name}",
        )
        destination.write_bytes(content)
        return self.import_file(
            destination,
            batch_uid=batch_uid,
            mapping_overrides=mapping_overrides,
        )

    def import_file(
        self,
        source: Path,
        *,
        batch_uid: str | None = None,
        mapping_overrides: dict[str, str] | None = None,
    ) -> CreatorImportBatch:
        managed_source = paths.resolve_under(paths.imports, source.name)
        if source.resolve() != managed_source:
            raise ImportValidationError("Import file must be in managed storage")
        relative_source = paths.relative(managed_source).as_posix()
        batch = self._imports.create_batch(
            batch_uid=batch_uid
            or f"IMP_{datetime.now(UTC):%Y%m%d}_{uuid4().hex[:12].upper()}",
            source_file=relative_source,
        )

        try:
            rows = self._read_rows(managed_source)
        except Exception as exc:
            batch.status = ImportStatus.FAILED
            batch.completed_at = datetime.now(UTC)
            batch.error_summary = str(exc)
            self._imports.add_error(
                batch_id=batch.id,
                row_number=0,
                raw_data={},
                error_type=type(exc).__name__,
                error_message=str(exc),
            )
            self._session.flush()
            return self._imports.get(batch.id) or batch

        batch.rows_total = len(rows)
        mapping = ColumnMapping(mapping_overrides)
        for row_number, raw_row in rows:
            try:
                with self._session.begin_nested():
                    parsed = self._parse_row(mapping.map_row(raw_row))
                    self._import_row(parsed)
                batch.rows_success += 1
            except Exception as exc:
                batch.rows_failed += 1
                self._imports.add_error(
                    batch_id=batch.id,
                    row_number=row_number,
                    raw_data=self._json_safe(raw_row),
                    error_type=type(exc).__name__,
                    error_message=str(exc),
                )

        batch.completed_at = datetime.now(UTC)
        if batch.rows_failed == 0:
            batch.status = ImportStatus.COMPLETED
        elif batch.rows_success == 0:
            batch.status = ImportStatus.FAILED
        else:
            batch.status = ImportStatus.PARTIAL
        if batch.rows_failed:
            batch.error_summary = (
                f"{batch.rows_failed} of {batch.rows_total} rows failed"
            )
        self._session.flush()
        return self._imports.get(batch.id) or batch

    def _parse_row(self, mapped: dict[str, Any]) -> ParsedImportRow:
        creator_values = {
            field: self._clean_scalar(mapped[field])
            for field in CREATOR_FIELDS
            if field in mapped
        }
        if "priority" in creator_values:
            try:
                creator_values["priority"] = int(creator_values["priority"])
            except (TypeError, ValueError) as exc:
                raise ImportValidationError("priority must be an integer") from exc
        if "active" in creator_values:
            creator_values["active"] = self._parse_bool(creator_values["active"])

        grouped_accounts: dict[Platform, dict[str, Any]] = {}
        for target, value in mapped.items():
            identity = self._platform_identity_target(target)
            if identity is None:
                continue
            platform, field = identity
            grouped_accounts.setdefault(platform, {})[field] = value

        generic_identity = {
            key: mapped.get(key)
            for key in ("platform", "username", "profile_url", "platform_user_id")
        }
        if any(generic_identity.values()):
            platform = self._normalizer.normalize_platform(
                generic_identity.pop("platform")
            )
            if platform is None:
                raise ImportValidationError(
                    "Generic account columns require a platform"
                )
            grouped_accounts.setdefault(platform, {}).update(
                {
                    key: value
                    for key, value in generic_identity.items()
                    if value is not None
                }
            )

        accounts = [
            self._normalizer.normalize_account(
                platform=platform,
                username=values.get("username"),
                profile_url=values.get("profile_url"),
                platform_user_id=values.get("platform_user_id"),
            )
            for platform, values in grouped_accounts.items()
        ]
        if not creator_values.get("display_name") and not (
            creator_values.get("creator_uid") or accounts
        ):
            raise ImportValidationError(
                "Row requires a creator name, creator_uid, or platform account"
            )
        return ParsedImportRow(creator_values=creator_values, accounts=accounts)

    def _import_row(self, row: ParsedImportRow) -> Creator:
        requested_uid = row.creator_values.get("creator_uid")
        uid_creator = (
            self._creators.get_by_uid(str(requested_uid)) if requested_uid else None
        )
        account_matches = {
            match.creator_id
            for account in row.accounts
            for match in self._accounts.find_matches(
                platform=account.platform,
                username=account.username,
                platform_user_id=account.platform_user_id,
                profile_url=account.profile_url,
            )
        }

        if uid_creator is not None:
            conflicting = account_matches - {uid_creator.id}
            if conflicting:
                duplicate_ids = conflicting | {uid_creator.id}
                self._creators.mark_possible_duplicates(duplicate_ids)
                raise ConflictError(
                    "creator_uid and platform accounts match different creators"
                )
            creator = uid_creator
        elif len(account_matches) == 1:
            creator = self._creators.get(account_matches.pop())
            if creator is None:
                raise ImportValidationError("Matched creator no longer exists")
        elif len(account_matches) > 1:
            self._creators.mark_possible_duplicates(account_matches)
            raise ConflictError(
                "Platform accounts match multiple creators; no merge was performed"
            )
        else:
            display_name = row.creator_values.get("display_name")
            if not display_name:
                raise ImportValidationError(
                    "A new creator requires display_name"
                )
            possible_matches = self._creators.possible_name_matches(
                str(display_name),
                row.creator_values.get("country"),
            )
            if possible_matches:
                self._creators.mark_possible_duplicates(
                    {candidate.id for candidate in possible_matches}
                )
                row.creator_values["possible_duplicate"] = True
            creator = self._creators.create(
                CreatorCreate(**row.creator_values).model_dump()
            )

        for account in row.accounts:
            matches = self._accounts.find_matches(
                platform=account.platform,
                username=account.username,
                platform_user_id=account.platform_user_id,
                profile_url=account.profile_url,
            )
            if any(match.creator_id != creator.id for match in matches):
                self._creators.mark_possible_duplicates(
                    {creator.id, *(match.creator_id for match in matches)}
                )
                raise ConflictError(
                    "Platform account belongs to another creator; no merge was performed"
                )
            if not matches:
                self._accounts.create(
                    creator.id,
                    {
                        "platform": account.platform,
                        "username": account.username,
                        "profile_url": account.profile_url,
                        "platform_user_id": account.platform_user_id,
                    },
                )
        return creator

    @staticmethod
    def _platform_identity_target(
        target: str,
    ) -> tuple[Platform, str] | None:
        suffixes = {
            "_username": "username",
            "_handle": "username",
            "_profile_url": "profile_url",
            "_platform_user_id": "platform_user_id",
        }
        for suffix, field in suffixes.items():
            if target.endswith(suffix):
                platform_name = target[: -len(suffix)]
                try:
                    return Platform(platform_name), field
                except ValueError as exc:
                    raise ImportValidationError(
                        f"Mapping targets unsupported platform: {platform_name}"
                    ) from exc
        return None

    @staticmethod
    def _read_rows(source: Path) -> list[tuple[int, dict[str, Any]]]:
        if source.suffix.casefold() == ".csv":
            return CreatorImportService._read_csv(source)
        if source.suffix.casefold() == ".xlsx":
            return CreatorImportService._read_xlsx(source)
        raise ImportValidationError("Only .csv and .xlsx files are supported")

    @staticmethod
    def _read_csv(source: Path) -> list[tuple[int, dict[str, Any]]]:
        raw = source.read_bytes()
        text = None
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ImportValidationError("CSV encoding is not supported")
        reader = csv.DictReader(StringIO(text))
        if not reader.fieldnames:
            raise ImportValidationError("CSV header row is missing")
        return [
            (row_number, dict(row))
            for row_number, row in enumerate(reader, start=2)
            if any(value is not None and str(value).strip() for value in row.values())
        ]

    @staticmethod
    def _read_xlsx(source: Path) -> list[tuple[int, dict[str, Any]]]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            try:
                headers = next(values)
            except StopIteration as exc:
                raise ImportValidationError("XLSX file is empty") from exc
            header_names = [
                str(value).strip() if value is not None else ""
                for value in headers
            ]
            if not any(header_names):
                raise ImportValidationError("XLSX header row is missing")
            rows = []
            for row_number, row_values in enumerate(values, start=2):
                if not any(value is not None and str(value).strip() for value in row_values):
                    continue
                rows.append(
                    (
                        row_number,
                        {
                            header: value
                            for header, value in zip(
                                header_names,
                                row_values,
                                strict=False,
                            )
                            if header
                        },
                    )
                )
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _parse_bool(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        normalized = str(value).strip().casefold()
        if normalized in {"1", "true", "yes", "y", "active"}:
            return True
        if normalized in {"0", "false", "no", "n", "inactive"}:
            return False
        raise ImportValidationError("active must be a boolean value")

    @staticmethod
    def _clean_scalar(value: Any) -> Any:
        return value.strip() if isinstance(value, str) else value

    @staticmethod
    def _json_safe(row: dict[str, Any]) -> dict[str, Any]:
        safe: dict[str, Any] = {}
        for key, value in row.items():
            if value is None or isinstance(value, (str, int, float, bool)):
                safe[str(key)] = value
            elif isinstance(value, (datetime, date)):
                safe[str(key)] = value.isoformat()
            else:
                safe[str(key)] = str(value)
        return safe
